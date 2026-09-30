#!/usr/bin/env python3
"""Build the editable Chinese submission-report draft from the user's DOCX."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("/Users/echomisty/Desktop/DASE7506 Project 1 TECH Report.docx")
OUTPUT = Path(__file__).resolve().parents[2] / "output" / "docx" / "DASE7506_Project_1_TECH_Report_中文初稿.docx"

BODY_CN = "Songti SC"
# LibreOffice on macOS does not consistently honor w:eastAsia for mixed runs.
# Using the same Chinese-capable face for all body-script slots keeps both
# Chinese and Latin text visible in Word and in the headless QA renderer.
BODY_LATIN = "Songti SC"
TITLE_LATIN = "Times New Roman"
HEADING_FONT = "Hiragino Sans GB"
NAVY = "1F4E78"
LIGHT_BLUE = "F3F7FB"
BORDER = "D9D9D9"


def set_run_font(run, *, latin: str = BODY_LATIN, east_asia: str = BODY_CN, size: float | None = None,
                 bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, *, latin: str, east_asia: str, size: float, bold: bool = False,
                   color: str = "000000") -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:cs"), latin)


def format_body(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(20)
    fmt.line_spacing = 1.08
    fmt.space_after = Pt(3)
    fmt.widow_control = True


def add_body(document: Document, text: str):
    paragraph = document.add_paragraph(text, style="Normal")
    format_body(paragraph)
    return paragraph


def add_heading(document: Document, text: str):
    paragraph = document.add_paragraph(text, style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    return paragraph


def add_caption(document: Document, text: str, *, page_break_before: bool = False):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run(text)
    set_run_font(run, size=9, bold=True, color="333333")
    return paragraph


def set_cell_margins(cell, top=70, start=95, bottom=70, end=95) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "5")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    tr_pr.append(marker)


def set_cell_width(cell, width_inches: float) -> None:
    width = Inches(width_inches)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.twips)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_inches: float) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(Inches(width_inches).twips)))
    tbl_w.set(qn("w:type"), "dxa")


def fill_cell(cell, text: str, *, header: bool = False, align: str = "left") -> None:
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    run = paragraph.add_run(text)
    set_run_font(run, size=8.7, bold=header, color="FFFFFF" if header else "222222")
    if header:
        shade_cell(cell, NAVY)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float],
              aligns: list[str] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, sum(widths))
    set_table_borders(table)
    aligns = aligns or ["left"] * len(headers)
    for col, text in enumerate(headers):
        set_cell_width(table.rows[0].cells[col], widths[col])
        fill_cell(table.rows[0].cells[col], text, header=True, align="center")
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col, text in enumerate(values):
            set_cell_width(row.cells[col], widths[col])
            fill_cell(row.cells[col], text, align=aligns[col])
            if row_index % 2 == 1:
                shade_cell(row.cells[col], LIGHT_BLUE)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bpb_equation(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(5)
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_para.append(math)

    prefix = OxmlElement("m:r")
    prefix_text = OxmlElement("m:t")
    prefix_text.text = "BPB = "
    prefix.append(prefix_text)
    math.append(prefix)

    fraction = OxmlElement("m:f")
    numerator = OxmlElement("m:num")
    numerator_run = OxmlElement("m:r")
    numerator_text = OxmlElement("m:t")
    numerator_text.text = "−Σᵢ ln p(yᵢ | y<ᵢ)"
    numerator_run.append(numerator_text)
    numerator.append(numerator_run)

    denominator = OxmlElement("m:den")
    denominator_run = OxmlElement("m:r")
    denominator_text = OxmlElement("m:t")
    denominator_text.text = "ln(2) × B"
    denominator_run.append(denominator_text)
    denominator.append(denominator_run)

    fraction.append(numerator)
    fraction.append(denominator)
    math.append(fraction)
    paragraph._p.append(math_para)


def remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def restore_preserved_package_parts(source: Path, output: Path) -> None:
    """Restore all source package parts except the intended document/style edits."""
    editable = {"word/document.xml", "word/styles.xml"}
    temp_output = output.with_suffix(".preserve-parts.docx")
    with ZipFile(source, "r") as source_zip, ZipFile(output, "r") as final_zip, ZipFile(temp_output, "w") as out_zip:
        source_names = set(source_zip.namelist())
        for info in final_zip.infolist():
            if info.filename in source_names and info.filename not in editable:
                data = source_zip.read(info.filename)
            else:
                data = final_zip.read(info.filename)
            out_zip.writestr(info, data)
    temp_output.replace(output)


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)

    # Preserve the source page geometry and establish portable fonts.
    set_style_font(document.styles["Normal"], latin=BODY_LATIN, east_asia=BODY_CN, size=10)
    set_style_font(document.styles["Title"], latin=TITLE_LATIN, east_asia=HEADING_FONT, size=22, color="000000")
    set_style_font(document.styles["Heading 1"], latin=HEADING_FONT, east_asia=HEADING_FONT, size=15.5, bold=True,
                   color="000000")
    document.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(4)

    paragraphs = document.paragraphs
    if len(paragraphs) < 6:
        raise RuntimeError("The source template no longer matches the expected six-paragraph structure")

    title, background_heading, background_first, task_heading, task_first, trailing_blank = paragraphs[:6]
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    background_heading.text = "1 背景"
    background_heading.style = document.styles["Heading 1"]

    background_first.text = (
        "语言模型通过前文估计下一个 token 的概率。模型对正确 token 分配的概率越高，编码同一段文本所需的信息量越少。"
        "因此，本项目关注的不只是模型能否拟合训练文本，还要检查它在未参与训练的数据上是否给出更准确、可复现的概率分布。"
    )
    format_body(background_first)

    paragraph = task_heading.insert_paragraph_before(
        "本项目使用课程提供的 WikiText-2 文本和固定 BPE-2048 tokenizer，从随机初始化训练 decoder-only Transformer。"
        "上下文长度固定为 256 tokens，评测按相互独立的因果窗口进行。模型只能利用当前位置之前的 token，不能读取未来信息，也不能在窗口之间保留隐藏状态。",
        style="Normal",
    )
    format_body(paragraph)
    task_heading.text = "2 任务目标与约束"
    task_heading.style = document.styles["Heading 1"]
    task_first.text = (
        "目标是在不使用外部文本、预训练权重或 test 调参的前提下，最小化完整 test split 的 bits per byte（BPB）。"
        "训练集用于更新参数；validation 用于选择模型结构、超参数、随机种子和 checkpoint；test 只在方案冻结后用于最终评测。"
    )
    format_body(task_first)
    remove_paragraph(trailing_blank)

    add_body(
        document,
        "数据文件、tokenizer 和课程评测器保持不变。评测时，每个窗口从空状态开始；除 split 的第一个 token 外，每个 target 恰好计入一次。"
        "这一规则保证不同模型面对相同的预测任务，也避免跨窗口信息带来额外优势。",
    )

    add_caption(document, "表 1 固定评测条件")
    add_table(
        document,
        ["项目", "固定设置", "作用"],
        [
            ["数据", "课程提供的 WikiText-2", "训练、验证和测试 split 的用途严格分开"],
            ["Tokenizer", "BPE-2048", "固定词表和分词过程"],
            ["上下文", "256 tokens", "训练与评测使用相同的最大上下文"],
            ["评测窗口", "独立且严格 causal", "禁止未来 token 和跨窗口状态"],
            ["模型选择", "仅使用 validation", "test 不参与结构或超参数选择"],
            ["最终评测", "CPU FP32", "报告可在 CPU 上复现的结果"],
        ],
        [1.15, 1.80, 2.72],
        ["center", "center", "left"],
    )

    add_heading(document, "3 评价指标")
    add_body(
        document,
        "主指标是 bits per byte。设 B 为该 split 原始 UTF-8 文本的字节数，求和覆盖所有被评分的 target。"
        "总负对数似然除以 ln(2) 后转换为比特数，再除以 B。",
    )
    add_bpb_equation(document)
    add_body(
        document,
        "BPB 越低越好。它以原始字节为分母，受 token 数量变化的影响小于 token-level perplexity。"
        "计算完整 split 时，应先累加所有 target 的负对数似然，再统一除以字节数，不能简单平均各 batch 或各窗口的 BPB。",
    )

    add_caption(document, "表 2 报告中的主要指标")
    add_table(
        document,
        ["指标", "定义或测量方式", "报告作用", "方向或上限"],
        [
            ["BPB", "总负对数似然 ÷ ln(2) ÷ UTF-8 字节数", "主要质量指标", "越低越好"],
            ["Token PPL", "exp（总 NLL ÷ target token 数）", "辅助解释 token 级预测难度", "越低越好"],
            ["CPU 时间比", "候选模型时间 ÷ 同机 baseline 时间", "衡量推理速度成本", "不超过 5×"],
            ["峰值 RSS", "评测进程的最大常驻内存", "衡量评测内存占用", "不超过 4 GiB"],
            ["推理资产", "推理所需资产的未压缩总大小", "控制 checkpoint 与依赖资源", "不超过 64 MiB"],
        ],
        [1.05, 2.20, 1.47, 0.95],
        ["center", "left", "left", "center"],
    )

    add_body(
        document,
        "Token PPL 适合说明模型在固定 tokenizer 下的 token 级不确定性，但它不是课程排名指标。"
        "资源指标也不与 BPB 合并成一个分数，而是作为硬约束单独检查；任何一项超出上限，都不能用更低的 BPB 抵消。",
    )

    add_heading(document, "4 评测与选择协议")
    add_body(
        document,
        "为避免 test 信息进入开发过程，模型开发和最终评测按四个阶段分开。每一阶段只使用与其目的相符的数据和信息。",
    )
    add_caption(document, "表 3 从训练到最终测试的流程")
    add_table(
        document,
        ["1 训练", "2 验证", "3 冻结", "4 测试"],
        [[
            "只用 train 更新参数",
            "用 validation 比较方案",
            "固定代码与 checkpoint",
            "一次性报告 test 指标",
        ]],
        [1.42, 1.42, 1.42, 1.42],
        ["center", "center", "center", "center"],
    )
    add_body(
        document,
        "不同方案的公平比较还需要固定评测器、上下文长度和已处理的训练 targets。"
        "当模型规模或训练时长发生变化时，报告应分别说明质量变化与计算代价，避免把更多训练或更大的模型误写成单一结构改进。",
    )

    # Reapply portable fonts to all visible runs, including those created by paragraph.text.
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else "Normal"
        for run in paragraph.runs:
            if style_name == "Title":
                set_run_font(run, latin=TITLE_LATIN, east_asia=HEADING_FONT)
            elif style_name == "Heading 1":
                set_run_font(run, latin=HEADING_FONT, east_asia=HEADING_FONT)
            else:
                set_run_font(run)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        # Preserve white, bold header formatting already applied.
                        current_color = run.font.color.rgb
                        current_bold = run.font.bold
                        set_run_font(run, size=8.7, bold=current_bold,
                                     color=str(current_color) if current_color is not None else None)

    document.save(OUTPUT)
    restore_preserved_package_parts(SOURCE, OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    try:
        build()
    except Exception as exc:
        print(f"build failed: {exc}", file=sys.stderr)
        raise
