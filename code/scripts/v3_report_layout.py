"""Shared Word layout primitives for the v3 report pair.

Extracted so the Chinese and English builders cannot drift apart: both call the
same table, caption, body and image helpers, and only the strings differ.

House rules enforced here:
  * tables are centred Morandi three-line tables (booktabs style);
  * every visible cell is centred, header rows included;
  * figure and table captions are centred with a bold "Fig. N." / "Table N." lead;
  * no caption is ever baked into an image.
"""
from __future__ import annotations

import struct
from pathlib import Path

from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

MORANDI_SLATE = '7E919D'
MORANDI_HEADER = 'E8ECE7'
MORANDI_ALT = 'F6F7F4'
TEXT = '252B2F'
CAPTION = '333333'

ZH_BODY = 'SimSun'
ZH_HEAD = 'SimHei'
EN_BODY = 'Times New Roman'
EN_HEAD = 'Times New Roman'


def png_size(path: Path) -> tuple[int, int]:
    """Read width and height from the PNG IHDR chunk without extra dependencies."""
    with path.open('rb') as handle:
        head = handle.read(24)
    if head[:8] != b'\x89PNG\r\n\x1a\n':
        raise ValueError(f'{path} is not a PNG')
    return struct.unpack('>II', head[16:24])


def set_font(run, *, latin: str, east_asia: str, size: float | None = None,
             bold: bool | None = None, color: str | None = None, italic: bool | None = None) -> None:
    run.font.name = latin
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:ascii'), latin)
    fonts.set(qn('w:hAnsi'), latin)
    fonts.set(qn('w:eastAsia'), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, *, latin: str, east_asia: str, size: float,
                   bold: bool = False, color: str = '000000') -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for slot in ('w:ascii', 'w:hAnsi', 'w:cs'):
        fonts.set(qn(slot), latin)
    fonts.set(qn('w:eastAsia'), east_asia)


def _border(parent, edge: str, *, val: str, size: int = 0, color: str = 'auto') -> None:
    node = parent.find(qn(f'w:{edge}'))
    if node is None:
        node = OxmlElement(f'w:{edge}')
        parent.append(node)
    node.set(qn('w:val'), val)
    node.set(qn('w:sz'), str(size))
    node.set(qn('w:space'), '0')
    node.set(qn('w:color'), color)


def _shade(cell, fill: str) -> None:
    pr = cell._tc.get_or_add_tcPr()
    node = pr.find(qn('w:shd'))
    if node is None:
        node = OxmlElement('w:shd')
        pr.append(node)
    node.set(qn('w:fill'), fill)
    node.set(qn('w:val'), 'clear')


def _cell_margins(cell, top=70, start=85, bottom=70, end=85) -> None:
    pr = cell._tc.get_or_add_tcPr()
    margins = pr.find(qn('w:tcMar'))
    if margins is None:
        margins = OxmlElement('w:tcMar')
        pr.append(margins)
    for name, value in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        node = margins.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def _repeat_header(row) -> None:
    pr = row._tr.get_or_add_trPr()
    if pr.find(qn('w:tblHeader')) is None:
        marker = OxmlElement('w:tblHeader')
        marker.set(qn('w:val'), 'true')
        pr.append(marker)


def _no_split(row) -> None:
    pr = row._tr.get_or_add_trPr()
    pr.append(OxmlElement('w:cantSplit'))


def three_line_table(document, headers: list[str], rows: list[list[str]],
                     widths: list[float], *, font: str, east_asia: str,
                     size: float = 8.5, header_size: float | None = None):
    """A centred Morandi three-line (booktabs) table with every cell centred."""
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement('w:tblW')
    tbl_w.set(qn('w:w'), str(int(Inches(total).twips)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_pr.append(tbl_w)

    borders = OxmlElement('w:tblBorders')
    tbl_pr.append(borders)
    _border(borders, 'top', val='single', size=12, color=MORANDI_SLATE)
    _border(borders, 'bottom', val='single', size=12, color=MORANDI_SLATE)
    for edge in ('left', 'right', 'insideH', 'insideV'):
        _border(borders, edge, val='nil')

    header_row = table.rows[0]
    _repeat_header(header_row)
    for column, text in enumerate(headers):
        width = OxmlElement('w:tcW')
        widths_for_cell = [widths[column]]
        cell = header_row.cells[column]
        cell.text = ''
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _cell_margins(cell)
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_pr.append(width)
        width.set(qn('w:w'), str(int(Inches(widths[column]).twips)))
        width.set(qn('w:type'), 'dxa')
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        set_font(run, latin=font, east_asia=east_asia,
                 size=header_size or size, bold=True, color=TEXT)
        _shade(cell, MORANDI_HEADER)
        tc_borders = OxmlElement('w:tcBorders')
        tc_pr.append(tc_borders)
        _border(tc_borders, 'bottom', val='single', size=7, color=MORANDI_SLATE)

    for row_index, values in enumerate(rows, start=1):
        row = table.add_row()
        _no_split(row)
        for column, text in enumerate(values):
            cell = row.cells[column]
            cell.text = ''
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = OxmlElement('w:tcW')
            tc_w.set(qn('w:w'), str(int(Inches(widths[column]).twips)))
            tc_w.set(qn('w:type'), 'dxa')
            tc_pr.append(tc_w)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt = paragraph.paragraph_format
            fmt.first_line_indent = Pt(0)
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.line_spacing = 1.0
            run = paragraph.add_run(text)
            set_font(run, latin=font, east_asia=east_asia, size=size,
                     bold=False, color=TEXT)
            _shade(cell, 'FFFFFF' if row_index % 2 else MORANDI_ALT)
    return table


def caption(document, label: str, text: str, *, font: str, east_asia: str,
            keep_with_next: bool = True, space_before: float = 4):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(5)
    fmt.keep_together = True
    fmt.keep_with_next = keep_with_next
    lead = paragraph.add_run(label)
    set_font(lead, latin=font, east_asia=east_asia, size=9, bold=True, color=CAPTION)
    body = paragraph.add_run(text)
    set_font(body, latin=font, east_asia=east_asia, size=9, color=CAPTION)
    return paragraph


def image(document, path: Path, width_inches: float):
    w, h = png_size(path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(2)
    fmt.space_after = Pt(2)
    fmt.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width_inches),
                    height=Emu(int(Inches(width_inches) * h / w)))
    return paragraph
