#!/usr/bin/env python3
"""Expand the reviewed Tech Report without overwriting the compact version.

Tables and prose remain native Word objects. Plot previews link conceptually to
the native chart deck and retained JSON; no experiment is run by this script.
"""
import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

import build_submission_report_docx as style

PROJECT = Path(__file__).resolve().parents[2]
SOURCE = PROJECT / "output/docx/DASE7506_Project_1_TECH_Report_中文修订版.docx"
OUTPUT = PROJECT / "output/docx/DASE7506_Project_1_TECH_Report_中文完整素材稿.docx"
FIGURES = PROJECT / "tmp/full_report_assets"
EVIDENCE = json.loads((PROJECT / "output/figures/full_report/evidence.json").read_text())
RUNS = EVIDENCE["runs"]
SUMMARY = json.loads((PROJECT / "code/results/2026-09-25/supplementary_summary.json").read_text())
RES = json.loads((PROJECT / "code/results/2026-09-23/resources_test.json").read_text())
document = Document(SOURCE)
figure_map = []
table_count = 0


def body(text):
    return style.add_body(document, text)


def heading(text):
    p = style.add_subheading(document, text)
    p.paragraph_format.page_break_before = True
    return p


def figure(key, explanation):
    record = next(f for f in EVIDENCE["figures"] if f["id"] == key)
    number = len(figure_map) + 1
    figure_map.append({"figure": f"E{number}", "asset": key, "ppt_slide": record["slide"]})
    style.add_figure(
        document, FIGURES / f"{key}.png",
        f"扩展图 E{number}  {record['title']}。{explanation}",
        alt_title=record["title"],
        alt_description=record["subtitle"] + "。" + record["note"],
    )
    return number


def table(title, headers, rows, widths=None):
    global table_count
    table_count += 1
    style.add_caption(document, f"扩展表 E{table_count}  {title}")
    return style.add_table(
        document, headers, rows,
        widths or [5.77 / len(headers)] * len(headers),
        ["left"] + ["center"] * (len(headers) - 1),
    )


def before(anchor, callback):
    """Append with valid package relationships, then move new body XML in place."""
    node = next(p._p for p in document.paragraphs if p.text == anchor)
    old = set(document._element.body)
    callback()
    new = [x for x in document._element.body if x not in old and x.tag != qn("w:sectPr")]
    for x in new:
        node.addprevious(x)


def b(run):
    return RUNS[run]["validation"]["bpb"]


def curve(run):
    r = RUNS[run]
    return [
        ((x.get("processed_targets") or x["step"] * r["args"]["batch_size"] * r["config"]["context"]), x["bpb"])
        for x in r["curve"]
    ]


def f(value, digits=6):
    return f"{value:.{digits}f}"


def status_edits():
    # Retain the actual assignment rule, but never call this long working draft compliant.
    all_paragraphs=list(document.paragraphs)+[p for t in document.tables for row in t.rows for cell in row.cells for p in cell.paragraphs]
    for p in all_paragraphs:
        for r in p.runs:
            revised=r.text.replace("0.006412","0.006411").replace("冻结清单记录哈希后，student-model test 只运行一次。","冻结清单记录哈希后，仅对该冻结配置作 test 评分；资源测量重复同一配置，没有用 test 选择新模型。")
            if revised!=r.text:r.text=revised
    for p in document.paragraphs:
        if p.text.startswith("同 seeds 17、137 的 RoPE 均值"):
            p.runs[0].text="在训练期 MPS 口径下，同 seeds 17、137 的 RoPE 均值为 1.594506，v1 为 1.627150，改善 0.032644。此处只比较两 seed 子集；RoPE 还改变了 checkpoint 选择，不能把全部差值归因于位置编码。"
            for r in p.runs[1:]:r.text=""
    for t in document.tables:
        for row in t.rows:
            if row.cells[0].text == "提交":
                style.fill_cell(row.cells[2], "本稿保留完整素材；提交前人工压缩至 10 页")
            if row.cells[0].text == "不超过 10 页的报告与 AI 披露":
                style.fill_cell(row.cells[1], "第 10 节已披露 AI；本素材稿需另行删减")
                style.fill_cell(row.cells[2], "待压缩", align="center")
    notice = style.add_body(
        document,
        "版本说明  本文件为中文完整素材稿，供人工修改与取舍，不按 10 页压缩。原 10 页修订版另行保留。正文保留概括性结论，扩展页补足曲线、数据和判断依据；附录收录全部 33 次已保存的非 smoke 实验及可选图形。提交前仍须满足课程页数要求。",
        first_line=False, bold_lead="版本说明  ",
    )
    first = next(p for p in document.paragraphs if p.text == "1 背景与任务目标")
    first._p.addprevious(notice._p)


def mainline():
    heading("4.1 优化路径的数值分解")
    figure("01_mainline", "各点是独立运行的结果；连线只帮助阅读开发顺序。")
    body("最初的问题是：在规模大致相同、训练 targets 相同的条件下，更换 block 是否值得继续投入。学生侧 LN+GELU 对照为 2.101695，RMSNorm+SwiGLU 为 2.040526，下降 0.061169 BPB。这个比较支持保留新 block，但不能直接说明 RMSNorm 和 SwiGLU 各自贡献多少，因此后续补做 2×2 消融。课程官方 baseline 为 2.071087，其训练设置不同，不能替代这里的受控对照。")
    body("随后将模型扩大到 w320 d8，短预算 BPB 降到 1.820432。较大的收益说明原小模型的容量可能限制了当前配方，因此先在新规模上选择学习率，再决定是否继续扩大。长训练进一步降至 1.622441，但这一步同时改变了训练预算与学习率 schedule，不能把 0.197991 的差值写成结构改进。")
    table("各阶段相邻差值及解释边界", ["相邻比较", "BPB 降低", "相对降低", "可以支持的判断"], [
        ["小型 block 更换", f(b("s0_lr1e3")-b("s_lr1e3")), "2.91%", "组合值得继续；需拆分机制"],
        ["小型至 w320 d8", f(b("s_lr1e3")-b("lr1e3_a")), "10.79%", "当前短预算下扩容有效"],
        ["w320 d8 短至长训练", f(b("lr1e3_a")-b("a_e8_lr1e3")), "10.88%", "增加训练投入有效"],
    ], [1.35, .85, .85, 2.72])
    body("决策顺序依次缩小了问题范围：先确认 block 组合是否有效，再找合适规模与学习率，最后检查更长训练的边际收益。它不是全因子搜索，也没有对所有组合和 seeds 穷举。")


def hyperparameters():
    heading("5.1 训练损失与 validation 为什么要同时看")
    figure("02_lr_loss", "各点为日志实际保留的 batch NLL；不同运行的记录步数不一定相同。")
    body("六个学习率的训练损失整体下降，但这些日志只记录部分 batch。0.0005 的记录位于 step 400、800、1200，其他五个运行位于 300、600、900、1200；图中按真实 step 放置点，不能把不同运行的第一个点视为相同进度。最终完整 validation 以 0.001 的 1.820432 最低。0.002 和 0.004 的记录损失较高，末尾 validation 也较差，当前预算下没有继续放大学习率的证据。")
    ids = ["lr5e4_a", "lr7p5e4_a", "lr1e3_a", "lr1p25e3_a", "lr2e3_a", "lr4e3_a"]
    table("六个学习率的末尾记录与完整验证", ["Peak LR", "末条 loss", "Validation BPB", "较 0.001 增量", "训练秒数"], [
        [str(RUNS[r]["args"]["lr"]), f(RUNS[r]["history"][-1]["loss"], 4), f(b(r)), f(b(r)-b("lr1e3_a")), f(RUNS[r]["train_seconds"], 1)] for r in ids
    ], [.85, 1.02, 1.2, 1.5, 1.2])
    body("因此后续 w320 d8 使用 0.001。0.00075 的结果为 1.826578，只落后 0.006146，现有单 seed 实验不足以证明二者存在稳定差距。这里选择观察值较好的设置，是有限预算下的操作决定。")
    heading("5.2 正则结果与学习率选择的适用范围")
    figure("04_regularization", "同一学习率与训练预算下的离散设置比较。")
    body("Dropout 0.05 使 validation 从 1.820432 升至 1.853883。增加 weight decay 到 0.2 后为 1.820517，与对照只差 0.000085。当前证据不支持通过增强这两项正则获得收益，因此保留 dropout 0 和 weight decay 0.1。不能把单次实验的近似持平解释为两种设置完全等价。")
    table("同一超参数在不同规模下的结果", ["规模", "LR 0.001", "LR 0.002", "本轮判断"], [
        ["小型 RMS+SwiGLU", f(b("s_lr1e3")), f(b("s_lr2e3")), "0.002 较好"],
        ["w320 d8", f(b("lr1e3_a")), f(b("lr2e3_a")), "0.001 较好"],
    ], [1.4, 1.1, 1.1, 2.17])
    body("小模型在 LR 0.002 下达到 1.953693，优于其 LR 0.001 结果。这条记录被保留，是因为它提醒我们：学习率结论依赖模型规模与预算。扩大模型后必须重新筛选，不能把小模型最优设置直接搬过去。主线图选择相同 LR 0.001 展示 block 和容量比较，而不是声称每个小模型点都经过充分调参。")
    heading("5.3 长训练曲线与停止扩张预算的理由")
    figure("05_v1_curves", "横轴按实际训练 targets 计算；旧日志中的 validation targets 不作为横轴。")
    rows1, rows2 = curve("a_e8_lr1e3"), curve("2026-09-23_seed137_final_recipe")
    table("v1 长配方在各验证点的 BPB", ["训练 targets", "seed 17", "seed 137", "seed 17 较前点降低"], [
        [f(x/1e6, 6)+"M", f(y), f(rows2[i][1]), "—" if not i else f(rows1[i-1][1]-y)]
        for i,(x,y) in enumerate(rows1)
    ], [1.4, 1.15, 1.15, 2.07])
    body("seed 17 在前两个验证点间下降约 0.164352，随后下降约 0.057641，末段只再下降 0.003166。边际收益明显缩小，因此当时没有立即投入 58M targets，而是用第二个 seed 检查配方。seed 137 的末尾为 1.631858，与 seed 17 相差 0.009417，两条曲线都支持长配方优于短配方。")
    body("最终按 validation 选择 seed 17，冻结后才评测 test。未运行 58M 只是当前预算分配决定，并不能证明更长训练或重新调度学习率一定无效。")


def final_results():
    heading("6.2 正式 test 改善对应什么")
    figure("06_test", "柱形从零开始，保持绝对 BPB 差值的可读性。")
    base, model = RES["baseline"], RES["models"][0]
    obs0, obs1 = base["observations"][0], model["observations"][0]
    table("test 总量与辅助指标", ["项目", "课程 baseline", "冻结 v1"], [
        ["BPB", f(base["bpb"]), f(model["bpb"])],
        ["Token PPL", f(base["token_ppl"], 3), f(model["token_ppl"], 3)],
        ["总 NLL（nats）", f(obs0["nll_nats"], 3), f(obs1["nll_nats"], 3)],
        ["被评分 targets", str(obs0["targets"]), str(obs1["targets"])],
        ["原始 UTF-8 bytes", str(obs0["utf8_bytes"]), str(obs1["utf8_bytes"])],
    ], [2.1, 1.835, 1.835])
    body("两个模型在相同 test 文本、相同 target 数与相同 tokenizer 上计分，因此 BPB 差异对应模型对同一文本的总预测负对数似然差异。相对改善为 (2.101265−1.646152)÷2.101265，约 21.66%。Token PPL 同步下降，但排名与主要结论仍按 BPB 表述。")
    body("这里没有加入 RoPE 的 test 柱。其 1.588095 是 validation 候选分数，不能与 v1 的 1.646152 test 分数直接相减，也不能据此宣布测试成绩进一步提升。")
    heading("6.3 资源门槛与测量协议")
    figure("07_resources", "两组均低于各自上限；标准化比例只用于判断合规。")
    candidate = SUMMARY["validation_resource_measurement"]["models"][0]
    table("v1 与 RoPE 候选各自的资源记录", ["项目", "v1 / test", "RoPE / validation"], [
        ["CPU 时间中位数（秒）", f(model["scoring_seconds_median"], 3), f(candidate["scoring_seconds_median"], 3)],
        ["3 次范围（秒）", f(model["scoring_seconds_min"],3)+"–"+f(model["scoring_seconds_max"],3), f(candidate["scoring_seconds_min"],3)+"–"+f(candidate["scoring_seconds_max"],3)],
        ["同协议 baseline 中位数（秒）", f(base["scoring_seconds_median"],3), f(SUMMARY["validation_resource_measurement"]["baseline"]["scoring_seconds_median"],3)],
        ["CPU 时间比 / 上限 5×", f(model["time_vs_baseline"],3)+"×", f(candidate["time_vs_baseline"],3)+"×"],
        ["峰值 RSS / 上限 4 GiB", f(model["peak_rss_gib_max"],3), f(candidate["peak_rss_gib_max"],3)],
        ["推理资产 / 上限 64 MiB", f(model["asset_mib"],3), f(candidate["asset_mib"],3)],
    ], [2.3, 1.735, 1.735])
    body("测量使用 CPU FP32、4 线程，预热后运行 3 次独立进程。时间取中位数，内存取各次峰值的最大值。v1 的正式资源检查与 RoPE 的候选检查分别使用 test 和 validation，硬件状态也可能不同，因此表中秒数不能用于推断更换位置编码造成了多少额外开销。")
    body("较低 BPB 只有在时间、内存和资产限制均满足时才可作为可用提交。本轮更宽的模型虽未超出 64 MiB，但其验证质量和剩余资源空间都不如 w320 d8，因此没有优先进入长训练。")


def ablations():
    heading("7.1 先排除训练器本身造成主要收益")
    figure("08_trainer", "四个点按设置变更顺序排列；均为小模型、seed 17、9.83M targets。")
    ids=["recipe_exact_baseline_s17","recipe_norepl_s17","recipe_norepl_beta095_s17","recipe_full_student_s17"]
    table("训练器设置变化的增量", ["设置", "Validation BPB", "相对前项变化"], [
        [name,f(b(r)),"—" if i==0 else f(b(r)-b(ids[i-1]))]
        for i,(name,r) in enumerate(zip(["复现 baseline 训练器","无放回采样","再改 beta2 为 0.95","完整学生训练设置"],ids))
    ], [2.5,1.5,1.77])
    body("无放回采样从 2.072376 变为 2.071598，只改善约 0.000778，接近持平。随后将 beta2 改为 0.95，BPB 升至 2.103114；完整学生训练设置为 2.098476，也未优于起点。因此不能把 v1 的主要收益归因于采样方式或训练器替换。")
    body("下一步把注意力放回结构：保持学生侧训练器，分别替换 norm 与 FFN。这里的 2.098476 是课程小模型结构下的训练器对照，和下一页学生侧 LN+GELU 的 2.101695 来自不同设置，不能混用。")
    heading("7.2 拆分 RMSNorm 与 SwiGLU 的贡献")
    figure("09_factorial", "两个 norm 与两个 FFN 构成 2×2 对照，均为同一短预算。")
    table("相对 LN+GELU 的改善", ["设置", "Validation BPB", "BPB 降低"], [
        [name,f(b(r)),f(b("s0_lr1e3")-b(r))] for name,r in [
            ("LN+GELU","s0_lr1e3"),("RMS+GELU","ablation_rms_gelu_s17"),
            ("LN+SwiGLU","ablation_ln_swiglu_s17"),("RMS+SwiGLU","s_lr1e3")]
    ], [2.4,1.6,1.77])
    body("SwiGLU 单独带来约 0.056130 BPB 的改善，RMSNorm 单独带来约 0.006411。组合改善为 0.061169，因此本轮 block 收益主要与 SwiGLU 替换相关。两项独立差值并不严格相加，不能把组合结果写成两个机制完全独立。")
    body("据此保留 RMSNorm+SwiGLU 作为后续基础结构，但不继续把单独更换 norm 当作主要搜索方向。该归因目前只有 seed 17 的证据，尚不能证明换 seed 后差值仍完全相同。")
    heading("7.3 更深、更宽为什么没有继续")
    figure("10_capacity", "质量与资产分别展示，避免把两种单位放在同一个纵轴。")
    ids=["lr1e3_a","capacity_w320d10_s17","capacity_w384d8_s17"]
    table("容量、质量与训练成本", ["设置", "参数量 M", "BPB", "资产 MiB", "训练秒数"], [
        [name,f(RUNS[r]["parameters"]/1e6,3),f(b(r)),f(RUNS[r]["asset_mib"],3),f(RUNS[r]["train_seconds"],1)]
        for name,r in zip(["w320 d8","w320 d10","w384 d8"],ids)
    ], [1.2,1.05,1.2,1.2,1.12])
    body("增加到 10 层后 BPB 为 1.872762，比 d8 高 0.052330；扩宽到 384 后为 1.829156，比原宽度高 0.008725。两者参数、资产与记录训练时间均增加，但当前短预算的质量没有提高，所以后续继续使用 w320 d8。")
    body("这个结果只否定了“在本轮相同 targets 与超参数下立即扩容”的优先级。更大模型可能需要不同学习率或更长训练，现有实验没有覆盖这些组合。为了避免把搜索预算用于尚未显示收益的方向，下一步选择改变位置表示。")
    heading("7.4 位置表示与输出权重的短实验")
    figure("11_representation", "四个点均为 seed 17；连接不同设置的线不表示连续训练。")
    ids=["lr1e3_a","ablation_untied_s17","ablation_rope_s17","ablation_rope_untied_s17"]
    table("位置与权重共享的同 seed 比较", ["设置", "BPB", "相对 learned tied 降低", "资产 MiB"], [
        [name,f(b(r)),f(b(ids[0])-b(r)),f(RUNS[r]["asset_mib"],3)]
        for name,r in zip(["Learned / tied","Learned / untied","RoPE / tied","RoPE / untied"],ids)
    ], [1.6,1.1,1.87,1.2])
    body("把 learned position 换成 RoPE，同时继续共享输入输出权重，BPB 从 1.820432 降至 1.754858，改善 0.065574。单独解除共享也有改善，但它需要更大的输出权重矩阵。RoPE tied 的资产约 40.033 MiB，略低于 learned tied 的 40.345 MiB，因此先把 RoPE 作为长训练候选方向。")
    body("同时改变位置表示和权重共享会混合两个因素，不能只看最低的 RoPE untied 点就宣称全部收益来自解除共享。下一步在 RoPE 条件下用两个相同 seeds 配对检查 untied 的额外贡献。")
    heading("7.5 untied 的平均额外收益是否值得投入")
    figure("12_rope_tying", "同一 seed 内比较 tied 与 untied，再汇总差值。")
    pairs=[("17","ablation_rope_s17","ablation_rope_untied_s17"),("137","ablation_rope_s137","ablation_rope_untied_s137")]
    table("RoPE tied / untied 配对", ["Seed", "Tied BPB", "Untied BPB", "额外降低"], [
        [seed,f(b(tied)),f(b(untied)),f(b(tied)-b(untied))] for seed,tied,untied in pairs
    ]+[
        ["均值",f(sum(b(x[1]) for x in pairs)/2),f(sum(b(x[2]) for x in pairs)/2),f(sum(b(x[1])-b(x[2]) for x in pairs)/2)]
    ], [.7,1.65,1.65,1.77])
    body("两个 seed 的 untied 都较好，平均额外改善为 0.006750 BPB，同时资产从约 40.033 增至 42.533 MiB。这个改善并非零，但低于本项目为分配长训练预算采用的 0.01 BPB 阈值。因此优先长训练结构较简单、资产较小的 RoPE tied。")
    body("0.01 是工程上的投入门槛，不是统计显著性标准；两个 seeds 也不足以得出稳定的总体效应。这里保留 untied 的结果，便于将来在预算允许时复核，而不把它写成失败机制。")
    heading("7.6 EMA 与 MTP 的对照口径")
    figure("13_training_addons", "两根柱使用不同但各自匹配的对照，不宜直接排序为通用方法优劣。")
    ema=RUNS["ablation_rope_untied_ema099_s17"]["averaging_comparison"]
    live=next(x["bpb"] for x in ema if x["weights"]=="live")
    averaged=next(x["bpb"] for x in ema if x["weights"]=="averaged")
    table("额外组件的实际比较", ["组件", "对照 BPB", "处理后 BPB", "额外降低"], [
        ["EMA 0.99，同 run live",f(live),f(averaged),f(live-averaged)],
        ["MTP k=1，匹配 untied 配方",f(b("ablation_rope_untied_s17")),f(b("ablation_rope_untied_mtp1_s17")),f(b("ablation_rope_untied_s17")-b("ablation_rope_untied_mtp1_s17"))],
    ], [2.1,1.22,1.22,1.23])
    body("EMA 需要和同一次运行中的 live 权重比较，收益才对应平均操作本身。同 run 的 live 为 1.735489，averaged 为 1.734143，实际差值约 0.001346；不能拿另一次 untied 运行的 1.745039 作对照，把运行间差异全部归给 EMA。")
    body("MTP 为 1.739964，相对匹配 untied 对照改善约 0.005075，也低于 0.01 阈值。因此本轮没有继续给 EMA 或 MTP 分配长训练预算。日志的单次训练秒数受运行顺序和系统负载影响，不能据此写出 MTP 固定增加某一比例的训练开销。")


def rope():
    heading("8.1 初次长训练的回退如何改变保存规则")
    figure("14_rope_initial", "初次长实验与后续 best-checkpoint 重跑分别记录，不能合并成同一次运行。")
    rows=curve("2026-09-25_rope_long_seed17")
    table("初次 RoPE 长训练的完整验证点", ["训练 targets", "MPS validation BPB", "权重保留情况"], [
        [f(x/1e6,6)+"M",f(y),"仅末尾权重保存" if i==len(rows)-1 else "保留了分数，未保留该点权重"]
        for i,(x,y) in enumerate(rows)
    ], [1.5,1.8,2.47])
    body("初次长训练在 19.66M targets 附近达到 1.620751，末尾 28.91M 却升到 1.634776，回退约 0.014025。由于当时只保存末尾 checkpoint，不能把中间验证分数当作已经可以提交的模型成绩。为验证和保留可用权重，后续进行了新的独立训练。")
    body("新运行预先声明 9.83M、19.66M、28.91M 三个完整验证点，每次改善都保存对应权重，训练仍完成全部搜索预算。这样把“训练实际投入多少”与“最终选出的权重训练到哪里”分开记录，避免把 28.91M 搜索成本误写为 19.66M。")
    heading("8.2 training loss 下降为何不能选末尾权重")
    figure("17_rope_training", "连接线仅连接真实日志点，最后相邻 step 的点保持实际距离。")
    body("三个 RoPE seed 的训练损失整体从约 3.7–3.8 降至约 2.3–2.4，但 validation 都在 19.66M 后回升。训练目标继续被拟合，并不保证对保留文本的预测继续改善。这种分离与过拟合或学习率、训练长度不匹配相容；现有实验没有单独区分两种原因，所以报告不把回退归结为唯一机制。")
    table("三个新运行的权重选择与真实投入", ["Seed", "选择 step", "权重 targets", "搜索 targets", "训练秒数"], [
        [str(r["seed"]),str(r["selected_step"]),f(r["selected_targets"]/1e6,2)+"M",f(r["search_targets"]/1e6,2)+"M",f(r["train_seconds"],1)]
        for r in SUMMARY["rope_long_training"]["runs"]
    ], [.7,1.05,1.3,1.3,1.42])
    body("下一步采用 validation 选择最佳权重，并在 CPU FP32 下重新计分，而不是按最后一个 training loss 的高低选择 seed。三次重跑都选中相同预算位置，说明这次规则调整对三个观察到的 seeds 都有实际作用；这并不保证换预算后最优位置仍为 19.66M。")
    heading("8.3 配对长方案比较与候选的边界")
    figure("16_paired_long", "只用共同 seeds 17、137；seed 233 另计入 RoPE 三 seed 统计。")
    rows=SUMMARY["rope_long_training"]["runs"]
    v1vals=[b("a_e8_lr1e3"),b("2026-09-23_seed137_final_recipe")]
    table("共同 seed 的长方案差值（均为 MPS）", ["Seed", "v1 末尾 BPB", "RoPE best BPB", "BPB 降低"], [
        [str(r["seed"]),f(v1vals[i]),f(r["validation_bpb"]),f(v1vals[i]-r["validation_bpb"])]
        for i,r in enumerate(rows[:2])
    ]+[["两 seed 均值",f(sum(v1vals)/2),f(sum(r["validation_bpb"] for r in rows[:2])/2),f(sum(v1vals)/2-sum(r["validation_bpb"] for r in rows[:2])/2)]],[.75,1.42,1.95,1.65])
    body("在共同的两个 seeds 中，RoPE 长方案平均比 v1 低约 0.032644 BPB。加入 seed 233 后，RoPE 三 seed 的均值为 1.600116、样本标准差为 0.011642。两 seed 配对均值与三 seed 总体描述回答不同问题，不能混用。")
    body("seed 137 的 CPU FP32 validation 为 1.588095，因此登记为后续候选。配对差值同时包含位置表示与 checkpoint 选择规则变化，不能全部解释为 RoPE 的单独贡献。候选已记录权重哈希并通过 validation 资源检查，但还没有正式 test 结果；本稿不自动改写冻结 v1 的提交成绩。")


GROUPS=[
    ("A.1 基线、学习率与正则",[
        ("R0_baseline","课程 baseline"),("s0_lr1e3","小型 LN+GELU"),("s_lr1e3","小型 RMS+SwiGLU / LR .001"),("s_lr2e3","小型 RMS+SwiGLU / LR .002"),
        ("lr5e4_a","w320 d8 / LR .0005"),("lr7p5e4_a","w320 d8 / LR .00075"),("lr1e3_a","w320 d8 / LR .001"),("lr1p25e3_a","w320 d8 / LR .00125"),("lr2e3_a","w320 d8 / LR .002"),("lr4e3_a","w320 d8 / LR .004"),("drop005_lr1e3_a","Dropout .05"),("wd2e1_lr1e3_a","Weight decay .2")]),
    ("A.2 训练器、block 与容量消融",[
        ("recipe_exact_baseline_s17","原始 baseline 设置"),("recipe_norepl_s17","无放回采样"),("recipe_norepl_beta095_s17","无放回 + beta2 .95"),("recipe_full_student_s17","完整学生训练设置"),
        ("ablation_rms_gelu_s17","RMS+GELU"),("ablation_ln_swiglu_s17","LN+SwiGLU"),("capacity_w320d10_s17","w320 d10"),("capacity_w384d8_s17","w384 d8")]),
    ("A.3 位置、共享权重与附加训练目标",[
        ("ablation_untied_s17","Learned untied"),("ablation_rope_s17","RoPE tied / seed 17"),("ablation_rope_s137","RoPE tied / seed 137"),("ablation_rope_untied_s17","RoPE untied / seed 17"),("ablation_rope_untied_s137","RoPE untied / seed 137"),("ablation_rope_untied_ema099_s17","RoPE untied + EMA .99"),("ablation_rope_untied_mtp1_s17","RoPE untied + MTP k1")]),
    ("A.4 长训练与最佳权重选择",[
        ("a_e8_lr1e3","v1 长配方 / seed 17"),("2026-09-23_seed137_final_recipe","v1 长配方 / seed 137"),("2026-09-25_rope_long_seed17","初次 RoPE 长训练"),("2026-09-25_rope_long_best_seed17","RoPE best / seed 17"),("2026-09-25_rope_long_best_seed137","RoPE best / seed 137"),("2026-09-25_rope_long_best_seed233","RoPE best / seed 233")]),
]


def appendices():
    expected={r for _,group in GROUPS for r,_ in group}
    assert expected == set(RUNS), f"Run inventory mismatch: {expected ^ set(RUNS)}"
    style.add_heading(document,"附录 A 全部已保存实验与来源索引",page_break=True)
    total=sum(r["train_seconds"] for r in RUNS.values())
    body(f"本附录覆盖 {len(RUNS)} 次非 smoke 运行，train_seconds 合计 {total:,.1f} 秒（{total/3600:.3f} 小时），不含验证、调试等时间。各组独立初始化。BPB 照录 metrics.json 的 validation，RoPE best 为 MPS 选择值，CPU FP32 值见表 9。EMA 行为平均权重；预算列表示实际搜索量。")
    counter=0
    for gi,(title,group) in enumerate(GROUPS):
        if gi: heading(title)
        else: style.add_subheading(document,title)
        rows=[]
        for run,name in group:
            counter+=1;r=RUNS[run]
            targets=r.get("processed_targets",r.get("train_tokens"))
            assert targets is not None, run
            assets=r.get("asset_mib",RES["baseline"]["asset_mib"] if run=="R0_baseline" else None)
            assert assets is not None, run
            rows.append([f"{counter:02d} {name}",str(r["seed"]),f(targets/1e6,2),f(b(run)),f(r["train_seconds"],1),f(r["parameters"]/1e6,3),f(assets,3)])
        table("保存的运行指标",["编号与设置","Seed","Targets M","Val. BPB","训练秒","参数 M","资产 MiB"],rows,[1.88,.45,.6,.8,.66,.66,.72])
        style.add_subheading(document,"对应 run 目录（位于 code/runs/）")
        first=counter-len(group)+1
        for i,(run,_) in enumerate(group,first):
            p=style.add_body(document,f"{i:02d}  {run}",first_line=False)
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.0
            for r in p.runs:style.base.set_run_font(r,size=8.7)
        if gi==3:
            body("三次 RoPE best 运行各实际处理 28.909568M targets，所选 step 2400 权重均对应 19.660800M targets。中间权重更优没有减少已经发生的训练搜索成本。初次 RoPE 长训练中间只留分数、未留权重，因此单独列出，不与新运行合并。")
    style.add_heading(document,"附录 B 半透明分层曲线图",page_break=True)
    body("以下是正文曲线的可选版式，不是额外实验。前后深度仅区分系列，纵轴截断；填充延伸至各层坐标底部，不是误差带。末点标签表示最后一次观测，不表示最优值。精确比较请同时看二维图和数据表。")
    wall_plots=json.loads((PROJECT/"output/figures/transparent_walls/data.json").read_text())
    for i,plot in enumerate(wall_plots):
        if i==2: heading("B.1 v1 与初次 RoPE 长训练")
        if i==4: heading("B.2 最新 RoPE 三 seed")
        style.add_figure(document,PROJECT/f"tmp/transparent_walls/{plot['id']}.png",f"图 B{i+1}  {plot['title']}。{plot['note']}",alt_title=plot['title'],alt_description="保存的实测点构成曲线，深度仅区分系列，透明填充不是误差带。"+plot['note'])
    style.add_heading(document,"附录 C 人工修改与图表来源",page_break=True)
    body("正文与表格均可在 Word 内直接修改。图 2 的合并流程图保留为独立 PPT，文字、方框和连接线可编辑。二维实验图放在“MP1_全部实验图表_可编辑.pptx”，数据系列、坐标轴、图例、说明均可编辑。修改 PPT 后导出图片，再在 Word 中更改对应图片。")
    rows=[["图 3","学习率 validation","3"],["图 4","RoPE 三 seed validation","15"]]
    rows += [[x["figure"],next(r["title"] for r in EVIDENCE["figures"] if r["id"]==x["asset"]),str(x["ppt_slide"])] for x in figure_map]
    table("Word 图与原生 PPT 页码对应",["Word 图号","实验内容","PPT 页"],rows,[1.0,3.87,.9])
    body("图 B1–B6 对应“MP1_半透明墙形曲线_可编辑.pptx”第 1–6 页。墙面、线、点和文字是独立原生对象，可修改透明度、颜色与标注；它们不是 Excel 联动图表。数值修改后需运行 build_transparent_wall_charts.mjs 重算几何位置，不宜手拖顶点代表新实验结果。")
    body("数据与绘图代码：output/figures/full_report/evidence.json、code/scripts/build_full_report_assets.mjs。墙形 SVG 与数据在 output/figures/transparent_walls/。Word 代码为 build_full_report_docx.py，流程图代码为 build_experiment_decision_flow.mjs，均在 code/scripts/。")
    body("手工修改后请另存，重跑脚本会覆盖生成内容。正式提交时可删去编辑说明和备选图、合并重复解释，但须保留机制对照、资源与冻结证据、AI 披露，并满足课程页数要求。")


status_edits()
before("5 超参数与训练长度的决策依据",mainline)
before("6 冻结后的最终结果与资源",hyperparameters)
before("7 补充消融如何改变后续方向",final_results)
before("8 RoPE 三 seed 与最佳 checkpoint",ablations)
before("9 正确性与提交要求",rope)
appendices()
document.core_properties.title = "DASE7506 Project 1 Tech Report — 中文完整素材稿"
document.core_properties.subject = "包含全部已保存实验、决策依据和可编辑图表来源的人工修改稿"
assert len(document.inline_shapes)==25, "Expected all original, added and alternate figures"
document.save(OUTPUT)
(PROJECT/"output/figures/full_report/word_figure_map.json").write_text(json.dumps(figure_map,ensure_ascii=False,indent=2))
print(OUTPUT)
print(f"Images: {len(document.inline_shapes)}; tables: {len(document.tables)}; runs: {len(RUNS)}")
