#!/usr/bin/env python3
"""Create an editable Chinese walkthrough and report-writing guide for MP1."""

from __future__ import annotations

import json
import statistics
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
CODE = ROOT / "code"
RESULTS = CODE / "results"
DATED = RESULTS / "2026-09-23"
ASSETS = RESULTS / "figures"
OUTPUT = ROOT / "output" / "docx" / "MP1_实验讲解与可编辑报告素材.docx"
NAVY = "17324D"
LIGHT_BLUE = "EAF1F8"
PALE_BLUE = "F4F7FA"
WHITE = "FFFFFF"
BLACK = "000000"
GRAY = "666666"
LIGHT_GRAY = "D9D9D9"
PENDING_PAGE_BREAK = False


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


freeze = load_json(DATED / "freeze_manifest.json")
test = load_json(DATED / "frozen_test_cpu_fp32.json")
resources = load_json(DATED / "resources_test.json")
seeds = load_json(DATED / "seed_comparison.json")


def load_run(name):
    return load_json(CODE / "runs" / name / "metrics.json")


def cpu_validation_bpb(item):
    name = Path(item["run_dir"]).name
    path = RESULTS / "2026-09-25" / f"{name}_validation_cpu_fp32.json"
    if path.exists():
        return load_json(path)["bpb"]
    return item["validation"]["bpb"]


SUPPLEMENTARY_RUNS = {
    name: load_run(name)
    for name in [
        "recipe_exact_baseline_s17", "recipe_norepl_s17",
        "recipe_norepl_beta095_s17", "recipe_full_student_s17",
        "s0_lr1e3", "ablation_rms_gelu_s17", "ablation_ln_swiglu_s17",
        "s_lr1e3", "lr1e3_a", "capacity_w320d10_s17",
        "capacity_w384d8_s17", "ablation_untied_s17",
        "ablation_rope_s17", "ablation_rope_s137",
        "ablation_rope_untied_s17", "ablation_rope_untied_s137",
        "ablation_rope_untied_ema099_s17", "ablation_rope_untied_mtp1_s17",
        "2026-09-25_rope_long_best_seed17",
        "2026-09-25_rope_long_best_seed137",
        "2026-09-25_rope_long_best_seed233",
    ]
}


def set_run_font(run, ascii_name="Heiti SC", east_asia="Heiti SC"):
    run.font.name = ascii_name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), ascii_name)
    fonts.set(qn("w:hAnsi"), ascii_name)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:hint"), "eastAsia")


def set_paragraph_keep(paragraph, keep_next=False, keep_lines=True):
    ppr = paragraph._p.get_or_add_pPr()
    if keep_next:
        ppr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        ppr.append(OxmlElement("w:keepLines"))


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tcmar = tcpr.first_child_found_in("w:tcMar")
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcmar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcmar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_GRAY, size="5"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def set_row_cant_split(row):
    """Keep each table row intact when Word lays it out across pages."""
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def set_cell_width(cell, width_inches):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(int(width_inches * 1440)))
    tcw.set(qn("w:type"), "dxa")


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_separate, text, fld_end])
    set_run_font(run)
    run.font.size = Pt(9)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    new_run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_text(paragraph, text, bold=False, italic=False, color=BLACK, size=None, font="Heiti SC"):
    run = paragraph.add_run(text)
    set_run_font(run, font)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_body(doc, text, bold_lead=None):
    para = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        add_text(para, bold_lead, bold=True)
        add_text(para, text[len(bold_lead):])
    else:
        add_text(para, text)
    return para


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    para = doc.add_paragraph(style=style)
    add_text(para, text)
    return para


def add_number(doc, text, level=0):
    style = "List Number" if level == 0 else "List Number 2"
    para = doc.add_paragraph(style=style)
    add_text(para, text)
    return para


def add_code(doc, text):
    para = doc.add_paragraph(style="Code")
    add_text(para, text, size=9, font="Courier New")
    return para


def add_heading(doc, text, level=1):
    global PENDING_PAGE_BREAK
    para = doc.add_paragraph(text, style=f"Heading {level}")
    if PENDING_PAGE_BREAK:
        para.paragraph_format.page_break_before = True
        PENDING_PAGE_BREAK = False
    for run in para.runs:
        set_run_font(run)
        run.font.color.rgb = RGBColor(0, 0, 0)
    set_paragraph_keep(para, keep_next=True)
    return para


def add_table(doc, headers, rows, widths, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_width(cell, widths[j])
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(header))
        set_run_font(run)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row_data in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        for j, value in enumerate(row_data):
            cell = row.cells[j]
            set_cell_width(cell, widths[j])
            set_cell_shading(cell, WHITE if i % 2 == 0 else PALE_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            if alignments:
                p.alignment = alignments[j]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_run_font(run)
            run.font.size = Pt(9.2)
            run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def set_picture_alt(inline_shape, title, description):
    docpr = inline_shape._inline.docPr
    docpr.set("title", title)
    docpr.set("descr", description)


def add_figure(doc, image_path, caption, source, width=7.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run)
    shape = run.add_picture(str(image_path), width=Inches(width))
    set_picture_alt(shape, caption, source)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(cap, caption, bold=True, size=9.5)
    src = doc.add_paragraph()
    src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    src.paragraph_format.space_after = Pt(9)
    add_text(src, source, italic=True, color=GRAY, size=8.5)
    return p


def add_page_break(doc):
    global PENDING_PAGE_BREAK
    PENDING_PAGE_BREAK = True


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Heiti SC"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Heiti SC")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Heiti SC")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    normal._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = "Heiti SC"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:ascii"), "Heiti SC")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Heiti SC")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    title._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Heiti SC"
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Heiti SC")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Heiti SC")
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    subtitle._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")

    heading_sizes = {1: 17, 2: 13, 3: 11}
    for level, size in heading_sizes.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Heiti SC"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Heiti SC")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Heiti SC")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
        style.paragraph_format.space_before = Pt(12 if level == 1 else 9)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    body = styles["Body Text"]
    body.font.name = "Heiti SC"
    body.font.size = Pt(10.5)
    body.font.color.rgb = RGBColor(0, 0, 0)
    body._element.rPr.rFonts.set(qn("w:ascii"), "Heiti SC")
    body._element.rPr.rFonts.set(qn("w:hAnsi"), "Heiti SC")
    body._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    body._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    body.paragraph_format.first_line_indent = Inches(0.25)
    body.paragraph_format.line_spacing = 1.32
    body.paragraph_format.space_after = Pt(6)

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = styles[name]
        style.font.name = "Heiti SC"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Heiti SC")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Heiti SC")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.2

    caption = styles["Caption"]
    caption.font.name = "Heiti SC"
    caption.font.size = Pt(9.5)
    caption.font.bold = True
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Heiti SC")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Heiti SC")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    caption._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")

    code = styles.add_style("Code", 1)
    code.font.name = "Courier New"
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor(0, 0, 0)
    code.paragraph_format.left_indent = Inches(0.25)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(5)
    code.paragraph_format.line_spacing = 1.05


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(hp, "Mini Project 1 实验讲解与可编辑报告素材", color=GRAY, size=8.5)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp)


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_document(doc)
    doc.core_properties.title = "Mini Project 1 实验讲解与可编辑报告素材"
    doc.core_properties.subject = "WikiText-2 语言模型实验过程 结果解释 图片来源与报告写作指南"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "MP1 WikiText-2 BPB Transformer 可编辑报告"

    # Cover
    doc.add_paragraph().paragraph_format.space_after = Pt(65)
    title = doc.add_paragraph("Mini Project 1 实验讲解与可编辑报告素材", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph("用于理解实验过程和独立撰写课程报告", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    add_body(doc, "这份文档解释项目从数据检查、模型设计、训练、validation 选择、模型冻结到最终 test 的完整过程，并补充冻结后的受控消融。文档不是要求你原样提交的成品答案，而是一份可以修改、删减和重新组织的实验材料。读完后，你应当能够用自己的语言回答三个问题：我们改了什么，为什么这样改，以及证据是否支持结论。")
    add_body(doc, f"最终结论是：冻结后的 seed 17 模型在 CPU FP32 test 上得到 {test['bpb']:.6f} BPB，课程基线为 {resources['baseline']['bpb']:.6f} BPB，相对降低约 21.66%。正式资源测量为基线时间的 {resources['models'][0]['time_vs_baseline']:.3f} 倍、峰值 RSS {resources['models'][0]['peak_rss_gib_max']:.3f} GiB、推理资产 {resources['models'][0]['asset_mib']:.3f} MiB，三个限制全部通过。冻结后实验发现 RoPE 在多个 seed 上进一步改善 validation，但它仍是 validation-only v2 候选，没有替换已冻结的 test 结论。")
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    add_table(
        doc,
        ["项目", "最终值", "应如何理解"],
        [
            ["Validation BPB", "1.622441", "用于模型和 seed 选择，不是最终 test 分数"],
            ["Test BPB", "1.646152", "冻结后运行一次得到的最终质量指标"],
            ["相对 test 改进", "21.66%", "相对于同一 test split 上的课程基线"],
            ["CPU 时间", "2.875x baseline", "低于 5x 上限"],
            ["峰值内存", "1.948 GiB", "低于 4 GiB 上限"],
            ["推理资产", "40.345 MiB", "低于 64 MiB 上限"],
            ["补充实验", "validation only", "与冻结 v1 test 结果严格分开"],
        ],
        [1.45, 1.65, 4.1],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(doc, "文档日期：2026 年 9 月 25 日。v1 实验结果冻结日期：2026 年 9 月 23 日。")
    add_page_break(doc)

    # Reading guide and contents
    add_heading(doc, "阅读方法", 1)
    add_body(doc, "如果你需要先快速理解项目，先阅读“任务和评测逻辑”“模型架构和数据流”“实验序列和证据链”“冻结测试和资源结果”四部分。如果你已经理解实验，只需要自己写报告，可以直接阅读“如何独立撰写最终报告”和“结论与证据对应表”。")
    add_heading(doc, "文档结构", 2)
    contents = [
        "第一部分 任务目标和评测逻辑",
        "第二部分 工作目录 数据和环境",
        "第三部分 模型架构和数据流",
        "第四部分 代码修复和实验可靠性",
        "第五部分 实验序列和证据链",
        "第六部分 冻结后补充消融",
        "第七部分 模型冻结 最终测试和资源限制",
        "第八部分 图片来源和制作过程",
        "第九部分 如何独立撰写最终报告",
        "第十部分 复现命令 文件位置和检查清单",
        "附录 指标词汇 哈希和常见问题",
    ]
    for item in contents:
        add_bullet(doc, item)
    add_heading(doc, "需要始终区分的三种结果", 2)
    add_table(
        doc,
        ["结果类型", "用途", "能否用于调参", "本项目实例"],
        [
            ["Train loss", "观察优化是否正常", "可以，但不能替代泛化指标", "训练日志中的交叉熵"],
            ["Validation BPB", "选择架构、超参数、seed 和 checkpoint", "可以", "seed 17 为 1.622441"],
            ["Test BPB", "冻结后报告最终结果", "不可以", "最终一次为 1.646152"],
        ],
        [1.4, 2.7, 1.4, 1.7],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(doc, "最重要的实验纪律是 test 不参与任何决策。因为一旦根据 test 改模型，test 就不再代表未知数据上的泛化能力。本项目先在 validation 上比较两个 seed，写入 freeze_manifest.json 后才运行最终 test。")
    add_page_break(doc)

    # Part 1
    add_heading(doc, "第一部分 任务目标和评测逻辑", 1)
    add_heading(doc, "任务要求", 2)
    add_body(doc, "项目要求从随机初始化训练一个因果语言模型，只允许使用提供的 WikiText-2 训练文本学习参数。固定 tokenizer、context 和 evaluator 不能因为追求分数而改变。模型需要在 CPU FP32 环境中可复现，并同时满足时间、内存和推理资产限制。")
    add_table(
        doc,
        ["固定要求", "具体值", "为什么固定"],
        [
            ["词表", "BPE-2048", "防止通过更换 tokenizer 获得不公平优势"],
            ["上下文", "256 tokens", "保证训练和评测窗口一致"],
            ["评测窗口", "独立且 causal", "防止跨窗口状态和未来信息泄漏"],
            ["主指标", "Bits per byte", "按原始字节归一化，便于公平比较"],
            ["时间限制", "不超过 baseline 的 5 倍", "约束 CPU 推理复杂度"],
            ["内存限制", "不超过 4 GiB", "约束实际部署占用"],
            ["资产限制", "不超过 64 MiB", "约束 checkpoint 和必要推理文件"],
        ],
        [1.5, 1.8, 3.9],
    )
    add_heading(doc, "Bits per byte 的含义", 2)
    add_body(doc, "对每个目标 token，模型产生一个条件概率。评测器提取正确 token 的负对数概率 NLL，将所有目标位置的 NLL 相加，再转换为 bit，并除以该 split 原始 UTF-8 字节数。可编辑的表达式如下。")
    add_code(doc, "BPB = total_nll_nats / ln(2) / raw_utf8_bytes")
    add_body(doc, "BPB 越低，表示平均每个原始字节需要的编码位数越少，即模型对文本的概率分配更准确。这个指标与 token 划分相比更加稳定，因此报告中应把 BPB 作为主指标，token perplexity 只作为补充。")
    add_heading(doc, "每个 target 如何被计数", 2)
    add_number(doc, "评测文本先由固定 tokenizer 转换成 token ID 序列。")
    add_number(doc, "序列被切成最多 257 个 token 的独立窗口。")
    add_number(doc, "窗口前 256 个 token 作为输入 x，后移一位的 256 个 token 作为目标 y。")
    add_number(doc, "模型只能看到当前位置之前的 token，attention mask 屏蔽未来位置。")
    add_number(doc, "最后不足 257 个 token 的短窗口仍然评分，避免 split 尾部被漏掉。")
    add_number(doc, "除 split 的第一个 token 外，每个 target 恰好计数一次。")
    add_body(doc, "契约测试专门检查 future-token invariance、batch independence、state reset、shifted loss gradient 和最后短窗口覆盖。五个测试全部通过。")
    # Part 2
    add_heading(doc, "第二部分 工作目录 数据和环境", 1)
    add_heading(doc, "项目目录", 2)
    add_body(doc, "所有实验、环境、模型、结果、图片和文档都位于 project1 文件夹内。项目根目录为 MP1_student_starter，代码位于 code，最终日期化结果位于 code/results/2026-09-23。")
    add_table(
        doc,
        ["路径", "作用", "是否需要修改"],
        [
            ["README.md", "快速了解结果 目录与复现入口", "克隆后首先阅读"],
            ["REQUIREMENTS.md", "需求、限制和完成状态", "通常不需要"],
            ["EXPERIMENT_REPORT_ZH.md", "逐步骤工程和实验记录", "可作为写作素材"],
            ["code/student.py", "最终学生模型实现", "核心提交文件"],
            ["code/train_student.py", "训练和 validation 流程", "记录训练 recipe"],
            ["code/runs", "每个实验的 metrics 和 curve", "不要手工修改结果"],
            ["code/results/2026-09-23", "冻结、test、资源和架构图", "最终证据目录"],
            ["output/pdf", "10 页正式 PDF", "成品参考"],
            ["output/docx", "本可编辑 Word 文档", "可以自由编辑"],
        ],
        [2.35, 3.3, 1.55],
    )
    add_heading(doc, "虚拟环境", 2)
    add_body(doc, "项目使用 code/.venv_project1，避免全局 Python 或其他课程项目的包版本影响实验。当前环境为 Python 3.13.12、PyTorch 2.7.1、NumPy 2.5.3 和 tokenizers 0.21.4。训练使用 Apple MPS BF16，最终质量和资源测量使用 CPU FP32。")
    add_body(doc, "VS Code 的 .vscode/settings.json 已把该环境设为默认解释器，终端默认进入 code 目录；.vscode/tasks.json 提供测试、结果表重建、图片生成、架构图生成和 test 资源测量任务。")
    add_heading(doc, "数据来源和完整性", 2)
    add_body(doc, "实际评分数据是项目 data 目录中的 WikiText-2 train、validation 和 test 文本，以及固定 tokenizer.json。数据来自课程 starter package，而不是额外爬取网页。PACKAGE_MANIFEST.json 和 data/manifest.json 中的 SHA-256 用于确认数据和固定评测器没有被意外改动。")
    add_table(
        doc,
        ["Split", "用途", "是否更新模型", "本项目使用方式"],
        [
            ["Train", "优化模型参数", "是", "随机窗口训练，处理 28,909,568 targets"],
            ["Validation", "选择模型和 seed", "否", "每次完整评分，最低 BPB 入选"],
            ["Test", "最终报告", "否", "冻结后完整运行一次"],
        ],
        [1.1, 2.2, 1.4, 2.5],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(doc, "训练文本经过 tokenizer 后约有 3,613,343 个 token。BatchLoader 每次读取 32 个窗口，每个窗口包含 257 个连续 token，因此每个 step 产生 32 × 256 = 8,192 个 next-token targets。最终 3,529 steps 对应约 28.91M processed targets。")
    # Part 3
    add_heading(doc, "第三部分 模型架构和数据流", 1)
    add_heading(doc, "最终模型配置", 2)
    add_table(
        doc,
        ["组件", "最终设置", "原理", "本项目中的作用"],
        [
            ["Backbone", "8 blocks width 320", "多层 decoder-only Transformer", "提高容量同时满足 CPU 和资产限制"],
            ["Attention", "4 heads head dim 80", "每个位置聚合此前 token 信息", "建模 256-token 依赖"],
            ["Normalization", "Pre-RMSNorm", "只根据均方根缩放激活", "简化计算并稳定 residual 优化"],
            ["Feed-forward", "SwiGLU 2.667x", "门控分支调节信息通过量", "提高相同规模下的表达效率"],
            ["Position", "learned embedding", "给不同 token 位置添加可学习向量", "适配固定 256 context"],
            ["Linear layers", "bias free", "减少冗余参数", "减小资产和计算量"],
            ["Output head", "tied weights", "输入 embedding 与输出矩阵共享", "节省参数并保持词空间一致"],
            ["Dropout", "0", "随机丢弃激活以正则化", "实验显示 0.05 反而变差"],
        ],
        [1.25, 1.5, 2.15, 2.3],
    )
    add_body(doc, "最终模型共有 10,570,560 个参数，checkpoint 为 40.345 MiB。它不是最大的可行模型，而是在验证性能、CPU 时间和资产限制之间选择的安全点。")
    add_figure(
        doc,
        ASSETS / "architecture_data_flow.png",
        "图 1 模型架构和端到端数据流",
        "来源：本项目 scripts/architecture_diagram.py 根据 student.py、train_student.py、evaluate.py、configs/ours_a.json 和冻结结果生成；未使用外部图片。",
        width=7.0,
    )
    add_heading(doc, "图 1 的逐层阅读", 2)
    add_number(doc, "Data preparation 展示训练文本、哈希检查、BPE-2048 分词、token ID、窗口批处理和 x y 平移。")
    add_number(doc, "Student model forward path 展示 embedding、8 个 Transformer blocks、RMSNorm、QKV、causal attention、SwiGLU 和 tied language-model head。")
    add_number(doc, "Optimization 展示交叉熵、反向传播、gradient clipping、AdamW、cosine learning-rate schedule 和 checkpoint。")
    add_number(doc, "Frozen evaluation 展示独立窗口、log probabilities、target NLL、BPB 聚合和日期化结果。")
    add_number(doc, "黄色流程专门表示 validation 选择、freeze manifest 和冻结后 test，目的是把 test 与调参隔离。")
    add_heading(doc, "单个 Transformer block 内部发生什么", 2)
    add_body(doc, "输入 hidden state 先经过 RMSNorm。线性层生成 Q、K、V，并拆分为 4 个 attention heads。causal scaled dot-product attention 只允许位置 t 访问不晚于 t 的 K 和 V。attention 结果经无 bias 投影后与输入做 residual addition。第二个 RMSNorm 后进入 SwiGLU feed-forward network，门控分支和 value 分支逐元素相乘，再投影回 width 320，并再次做 residual addition。这个过程重复 8 次。")
    add_body(doc, "最终 RMSNorm 的输出与 token embedding 权重相乘得到 2,048 个词表 logits，再经 log-softmax 变成 log probabilities。训练时提取目标 token 的负对数概率形成 cross-entropy；评测时将全部 target NLL 汇总为 BPB。")
    # Part 4
    add_heading(doc, "第四部分 代码修复和实验可靠性", 1)
    add_heading(doc, "RoPE 广播维度修复", 2)
    add_body(doc, "Attention 的 Q 和 K 形状为 batch heads time head_dim。RoPE 的 cosine 和 sine 必须在 batch 与 heads 两个维度上广播，因此形状应该是 1 1 time head_dim。原实现是 1 time 1 head_dim，会把时间维错误地对齐到 head 维。")
    add_code(doc, "错误形式  emb.cos()[None, :, None, :]\n正确形式  emb.cos()[None, None, :, :]")
    add_body(doc, "本项目最终配置使用 learned positional embedding，但修复 RoPE 仍然必要，因为 student.py 支持可选 RoPE。修复后模型输出形状正确，log probability normalization 误差约 4.77e-7，远低于测试允许值。")
    add_heading(doc, "训练曲线横轴修复", 2)
    add_body(doc, "学习曲线需要用累计 processed training targets 作为横轴。原代码先保存训练 targets，随后 validate 返回的固定 validation target 数 376,599 覆盖了同名字段，导致所有点横坐标相同。修复后分别记录 processed_targets 和 validation_targets。旧 run 通过 step × batch_size × 256 恢复真实横轴，不需要重新训练。")
    add_heading(doc, "资源测量脚本修复", 2)
    add_body(doc, "资源脚本原来包含语法错误并硬编码旧环境与旧基线时间。新版本使用当前 sys.executable，每个 repetition 启动全新进程，固定 4 个 CPU 线程，对 baseline 和 candidate 使用相同 split 与 FP32，先预热 1 次，再正式运行 3 次。时间取中位数，峰值 RSS 取最大值。")
    add_heading(doc, "实验汇总和绘图修复", 2)
    add_body(doc, "collect.py 现在从每个正式 run 的 metrics.json 重建 grid.tsv，固定字段顺序并排除 smoke run。plot.py 直接读取结构化结果生成 SVG，不依赖 GUI 绘图工具。architecture_diagram.py 根据当前代码、配置和冻结结果生成架构图。")
    add_heading(doc, "正确性测试", 2)
    add_table(
        doc,
        ["测试", "检查内容", "结果"],
        [
            ["Future input invariance", "修改未来 token 不应改变更早预测", "通过"],
            ["Probability normalization", "每个位置概率总和为 1", "通过"],
            ["Batch independence", "一个样本不能影响另一个样本", "通过"],
            ["Window state reset", "窗口之间不保留隐状态", "通过"],
            ["Shifted gradient", "next-token loss 能产生有限非零梯度", "通过"],
            ["Final short window", "最后不足 context 的 target 不遗漏", "通过"],
        ],
        [2.0, 4.1, 1.1],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_body(doc, "最后短窗口检查与其他检查合并在五个 unittest 方法中，因此测试输出为 5 tests，而不是 6 个独立方法。")
    # Part 5
    add_heading(doc, "第五部分 实验序列和证据链", 1)
    add_body(doc, "实验不是一次直接猜中最终配置，而是按问题逐层缩小范围。公平比较的核心是改变一个主要因素时尽量固定其他条件。短实验统一处理 9,830,400 targets，batch size 32、context 256、seed 17、AdamW 和 validation evaluator 保持一致。")
    add_heading(doc, "实验零 官方基线", 2)
    add_body(doc, "先运行课程 baseline，目的是确认数据、tokenizer、evaluator 和 CPU 测量链路正常，并给资源比值提供同机参照。baseline 有 1.088M 参数，validation BPB 为 2.071087，test BPB 为 2.101265。")
    add_code(doc, "python train.py --config configs/baseline.json --run_dir runs/R0_baseline ...")
    add_body(doc, "报告中 baseline 不是为了证明 student trainer 一定更好，而是确定一个可复现起点。模型改进应同时与 baseline 和受控消融进行比较。")
    add_heading(doc, "实验一 高效 block 消融", 2)
    add_table(
        doc,
        ["Run", "Norm", "FFN", "Targets", "Validation BPB"],
        [
            ["s0_lr1e3", "LayerNorm", "GELU", "9.83M", "2.101695"],
            ["ablation_rms_gelu_s17", "RMSNorm", "GELU", "9.83M", "2.095283"],
            ["ablation_ln_swiglu_s17", "LayerNorm", "SwiGLU", "9.83M", "2.045564"],
            ["s_lr1e3", "RMSNorm", "SwiGLU", "9.83M", "2.040526"],
        ],
        [2.05, 1.3, 1.2, 1.15, 1.5],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_body(doc, "四个模型的参数量和训练量接近。固定 GELU 时，RMSNorm 改善 0.006412 BPB；固定 LayerNorm 时，SwiGLU 改善 0.056130 BPB；组合改善 0.061169 BPB。由此可把主要收益归因于 SwiGLU，RMSNorm 提供较小的附加收益，两者交互接近零。")
    add_heading(doc, "实验二 扩大模型容量", 2)
    add_table(
        doc,
        ["Run", "模型规模", "Targets", "Validation BPB", "差异"],
        [
            ["s_lr1e3", "1.082M", "9.83M", "2.040526", "小型高效 block"],
            ["lr1e3_a", "10.571M", "9.83M", "1.820432", "w320 d8"],
        ],
        [1.25, 1.4, 1.25, 1.45, 1.85],
    )
    add_body(doc, "在相同 processed targets 下，大模型比小型 RMSNorm + SwiGLU 模型降低 0.220094 BPB，约 10.79%。这说明项目在 1M 参数附近明显欠拟合，扩大有效容量比微调 regularization 更重要。这里同时改变了 width 和 depth，所以结论应写成“扩大整体容量有效”，不能把收益单独归因于加宽或加深。")
    add_heading(doc, "实验三 Peak learning rate 搜索", 2)
    add_table(
        doc,
        ["Peak LR", "Validation BPB", "与最佳差值", "解释"],
        [
            ["5e-4", "1.887329", "+0.066897", "更新偏慢"],
            ["7.5e-4", "1.826578", "+0.006146", "接近最佳"],
            ["1e-3", "1.820432", "0", "最终选择"],
            ["1.25e-3", "1.850677", "+0.030245", "开始恶化"],
            ["2e-3", "1.995888", "+0.175456", "明显过高"],
            ["4e-3", "1.972859", "+0.152427", "明显过高"],
        ],
        [1.2, 1.65, 1.55, 2.8],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(doc, "曲线在 1e-3 附近最低，7.5e-4 很接近，但 1.25e-3 以上明显变差。因此停止扩展更高学习率，并将 1e-3 固定到后续长训练。")
    add_heading(doc, "实验四 Dropout 和 weight decay", 2)
    add_table(
        doc,
        ["因素", "默认值", "实验值", "默认 BPB", "实验 BPB", "结论"],
        [
            ["Dropout", "0", "0.05", "1.820432", "1.853883", "0.05 变差 0.033451"],
            ["Weight decay", "0.1", "0.2", "1.820432", "1.820517", "差异仅 0.000085"],
        ],
        [1.15, 1.0, 1.0, 1.25, 1.25, 1.55],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 4 + [WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(doc, "Dropout 0.05 使结果明显恶化，说明在 9.83M targets 阶段模型仍然更接近欠拟合，而不是需要更强随机正则化。weight decay 从 0.1 提高到 0.2 几乎没有变化，因此保留较简单的 0.1。负面结果仍然应该写进报告，因为它们说明哪些方向已经被验证且不值得继续消耗预算。")
    add_heading(doc, "实验五 延长训练", 2)
    add_table(
        doc,
        ["Run", "模型", "Targets", "Validation BPB", "相对变化"],
        [
            ["lr1e3_a", "w320 d8", "9.83M", "1.820432", "短训练"],
            ["a_e8_lr1e3", "w320 d8", "28.91M", "1.622441", "降低约 10.88%"],
        ],
        [1.4, 1.45, 1.25, 1.5, 1.6],
    )
    add_body(doc, "模型和 peak LR 保持不变，只延长 schedule。Validation BPB 从 1.820432 降至 1.622441，说明增加数据遍历次数仍然有显著收益。曲线后段逐渐变平，为是否继续 58M targets 提供依据。")
    add_heading(doc, "实验六 第二个 seed 和可迁移性", 2)
    add_table(
        doc,
        ["Seed", "Targets", "Validation BPB", "Checkpoint hash 前 12 位", "选择"],
        [
            ["17", "28,909,568", "1.622441", "ab9ba8b64851", "冻结"],
            ["137", "28,909,568", "1.631858", "48cfecb3ed92", "未选择"],
        ],
        [0.8, 1.45, 1.5, 2.35, 1.1],
        [WD_ALIGN_PARAGRAPH.CENTER] * 5,
    )
    add_body(doc, "两个 seed 相差 0.009417 BPB，但都显著优于相同模型的 9.83M-target 结果。它支持“长训练收益可以迁移到另一个随机初始化”，但两个 seed 仍不足以准确估计均值和标准差。报告中应写“提供稳定性证据”，不要写“已经证明对所有 seed 稳定”。")
    add_heading(doc, "为什么没有运行 58M targets", 2)
    add_body(doc, "Seed 17 在最后约 7.2M targets 只从 1.625607 改善到 1.622441，约 0.003166 BPB。Seed 137 最终又比 seed 17 差 0.009417 BPB。根据这两个现象，继续把训练量翻倍的预期收益已经很小，而且无法保证超过 seed 波动。因此选择停止在 28.91M targets。这是基于 validation 曲线和成本收益作出的决定，与 test 结果无关。")
    add_figure(
        doc,
        ASSETS / "experiment_summary.png",
        "图 2 模型质量 优化和资源效率汇总",
        "来源：本项目 scripts/plot.py 读取 runs/*/metrics.json、curve.json、results/grid.tsv 和资源测量 JSON 生成；未使用外部图片或外部实验数据。",
        width=7.0,
    )
    add_heading(doc, "图 2 四个面板如何解释", 2)
    add_bullet(doc, "面板 a 学习动态：比较小型 LN + GELU、小型 RMSNorm + SwiGLU、大型模型和长训练。重点是容量与训练长度带来的阶梯式改进。")
    add_bullet(doc, "面板 b 学习率搜索：横轴为 peak LR，纵轴为 validation BPB。1e-3 被圈出，是相同预算下最低点。")
    add_bullet(doc, "面板 c 质量和容量：横轴为参数量，纵轴为 BPB，点大小表示 processed targets。它同时显示增加容量和增加训练量的收益。")
    add_bullet(doc, "面板 d 资源余量：把 CPU、RAM 和资产转换为各自上限的百分比。资产最接近限制，约占 63%。")
    # Part 6 supplementary validation-only ablations
    r = SUPPLEMENTARY_RUNS
    add_heading(doc, "第六部分 冻结后补充消融", 1)
    add_body(doc, "这一部分的实验在 v1 test 已经冻结并运行之后完成，目的不是回头修改 v1，而是补齐因果证据并形成下一轮候选。所有选择仍只使用 validation。报告必须把这些数字标为 validation-only，不能把它们写成新的 test 成绩。")
    add_heading(doc, "训练配方是否解释了模型收益", 2)
    add_table(
        doc,
        ["设置", "Sampling", "Schedule", "Beta2", "Validation BPB"],
        [
            ["课堂基准配方", "有放回", "baseline", "0.999", "2.072376"],
            ["仅改无放回", "无放回", "baseline", "0.999", "2.071598"],
            ["再改 beta2", "无放回", "baseline", "0.95", "2.103114"],
            ["完整学生配方", "无放回", "cosine", "0.95", "2.098476"],
        ],
        [1.65, 1.25, 1.35, 1.0, 1.55],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 4,
    )
    add_body(doc, "无放回采样仅改善 0.000778 BPB，可视为中性；beta2 从 0.999 改为 0.95 后反而恶化 0.031516 BPB，cosine schedule 只回收其中一小部分。这个结果排除了“最终收益只是 trainer 更强”的解释。")
    add_heading(doc, "容量形状筛选", 2)
    add_table(
        doc,
        ["配置", "参数量", "资产", "Targets", "Validation BPB", "决定"],
        [
            ["w320 d8", "10.571M", "40.345 MiB", "9.83M", "1.820432", "保留"],
            ["w320 d10", "13.029M", "49.728 MiB", "9.83M", "1.872762", "拒绝"],
            ["w384 d8", "15.047M", "57.422 MiB", "9.83M", "1.829156", "拒绝"],
        ],
        [1.05, 1.15, 1.25, 1.05, 1.45, 1.0],
        [WD_ALIGN_PARAGRAPH.CENTER] * 6,
    )
    add_body(doc, "更深的 w320 d10 明显变差；更宽的 w384 d8 也没有超过控制组，而且资产已达到 57.422 MiB，距离 64 MiB 上限很近。因此没有继续运行 d12，也没有把宽度扩到 384 以上。")
    add_heading(doc, "为什么这些负面结果重要", 2)
    add_body(doc, "如果只展示最终成功模型，读者无法判断容量、训练器或表示层分别贡献了什么。配方分解说明架构收益不是由采样或优化器参数伪造；容量筛选说明参数更多并不自动更好；这两组负面结果共同支持 w320 d8 是当前预算下的合理形状。")
    add_heading(doc, "位置编码和权重共享", 2)
    add_table(
        doc,
        ["Position", "Head", "Seed", "Validation BPB", "资产"],
        [
            ["learned", "tied", "17", "1.820432", "40.345 MiB"],
            ["learned", "untied", "17", "1.786678", "42.846 MiB"],
            ["RoPE", "tied", "17", "1.754858", "40.033 MiB"],
            ["RoPE", "tied", "137", "1.741184", "40.033 MiB"],
            ["RoPE", "untied", "17", "1.745039", "42.511 MiB"],
            ["RoPE", "untied", "137", "1.737502", "42.511 MiB"],
        ],
        [1.5, 1.35, 0.8, 1.65, 1.45],
        [WD_ALIGN_PARAGRAPH.CENTER] * 5,
    )
    rope_tied = [r["ablation_rope_s17"]["validation"]["bpb"], r["ablation_rope_s137"]["validation"]["bpb"]]
    rope_untied = [r["ablation_rope_untied_s17"]["validation"]["bpb"], r["ablation_rope_untied_s137"]["validation"]["bpb"]]
    add_body(doc, f"RoPE tied 的双 seed 均值为 {statistics.mean(rope_tied):.6f}，RoPE untied 为 {statistics.mean(rope_untied):.6f}。解开输出权重仅额外改善 {statistics.mean(rope_tied) - statistics.mean(rope_untied):.6f} BPB，低于消融协议预先设定的 0.01 阈值。为了保留更小资产和更简单实现，长训练使用 RoPE tied。")
    add_heading(doc, "EMA 和多步预测", 2)
    add_table(
        doc,
        ["机制", "基准", "Validation BPB", "观察", "决定"],
        [
            ["EMA 0.99", "RoPE untied", "1.734143", "比同 run live 权重仅好约 0.00135", "拒绝"],
            ["MTP k 1", "RoPE untied", "1.739964", "未超过基准且训练约慢 12%", "拒绝"],
        ],
        [1.2, 1.45, 1.4, 2.45, 1.0],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_body(doc, "EMA 的改善远小于 seed 波动，MTP 没有获得质量收益。二者都会让训练和解释更复杂，因此没有进入长训练。")
    add_heading(doc, "每个方向的结论", 2)
    add_table(
        doc,
        ["方向", "结果", "是否继续"],
        [
            ["Sampling", "几乎中性", "不作为主要贡献"],
            ["Beta2 0.95", "小 baseline 上变差", "只保留在已验证的最终配方"],
            ["RMSNorm", "小幅稳定收益", "保留"],
            ["SwiGLU", "主要 block 收益", "保留"],
            ["更深或更宽", "无改善且更贵", "停止"],
            ["RoPE", "两个 seed 都改善", "进入长训练"],
            ["Untied head", "增益低于阈值", "停止"],
            ["EMA 和 MTP", "收益不足", "停止"],
        ],
        [1.75, 3.0, 2.45],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_heading(doc, "长训练和最佳 validation 检查点", 2)
    add_body(doc, "第一条 RoPE seed 17 长训练只保存了末尾权重。它在 19.66M targets 达到 1.620751，但最终 28.91M targets 回退到 1.634776。这说明固定保存最后一步会丢失最佳 validation 模型。训练脚本随后增加 eval-at-targets 和 select-best-validation 两个显式控制。")
    long_runs = [
        r["2026-09-25_rope_long_best_seed17"],
        r["2026-09-25_rope_long_best_seed137"],
        r["2026-09-25_rope_long_best_seed233"],
    ]
    add_table(
        doc,
        ["Seed", "Search targets", "Selected targets", "CPU FP32 BPB", "Final MPS BPB", "Step"],
        [
            [str(item["seed"]), f"{item['processed_targets']:,}", f"{item['checkpoint_processed_targets']:,}", f"{cpu_validation_bpb(item):.6f}", f"{item['curve'][-1]['bpb']:.6f}", str(item["selected_step"])]
            for item in long_runs
        ],
        [0.7, 1.4, 1.45, 1.35, 1.25, 0.9],
        [WD_ALIGN_PARAGRAPH.CENTER] * 6,
    )
    long_bpbs = [cpu_validation_bpb(item) for item in long_runs]
    add_body(doc, f"三个 seed 的 selected BPB 均值为 {statistics.mean(long_bpbs):.6f}，样本标准差为 {statistics.stdev(long_bpbs):.6f}。三个 run 都实际搜索到 28.91M targets，但被保存的最佳权重位于 19.66M targets。报告应同时写 search targets 和 selected targets，避免把搜索成本写小。")
    add_body(doc, "RoPE 长训练结果比 learned-position v1 的 validation 更低，但没有运行新的 test。若课程允许提交 v2，应先根据 CPU FP32 validation 选定唯一 checkpoint，写入新的 freeze manifest，再运行一次 test 和资源测量。")
    add_code(doc, "python train_student.py --config configs/ours_a_rope.json --epochs 8 --eval-at-targets 9830400,19660800,28909568 --select-best-validation ...")
    # Part 7
    add_heading(doc, "第七部分 模型冻结 最终测试和资源限制", 1)
    add_heading(doc, "冻结的含义", 2)
    add_body(doc, "冻结不是简单地停止训练，而是把最终 checkpoint、student.py、evaluate.py、tokenizer.json、模型配置、seed、训练 targets 和 validation BPB 一起记录。冻结以后不能因为 test 结果不理想而换 seed、调超参数或继续训练。")
    add_table(
        doc,
        ["对象", "SHA-256"],
        [
            ["Final checkpoint", freeze["checkpoint_sha256"]],
            ["student.py", freeze["implementation_sha256"]],
            ["evaluate.py", freeze["evaluator_sha256"]],
            ["tokenizer.json", freeze["tokenizer_sha256"]],
        ],
        [1.55, 5.65],
    )
    add_heading(doc, "冻结后 test 结果", 2)
    add_table(
        doc,
        ["模型", "Test BPB", "Token PPL", "NLL nats", "Targets", "UTF-8 bytes"],
        [
            ["Baseline", "2.101265", "80.8479", "1,881,798.91", "428,405", "1,292,013"],
            ["Final model", "1.646152", "31.2239", "1,474,220.05", "428,405", "1,292,013"],
        ],
        [1.15, 1.05, 1.05, 1.55, 1.2, 1.3],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 5,
    )
    add_body(doc, "绝对 BPB 降低为 2.101265 - 1.646152 = 0.455113。相对改进约为 0.455113 / 2.101265 = 21.66%。Validation 与 test 的相对改进恰好都约为 21.66%，但它们使用不同 split 上各自的 baseline，报告中不能混用原始值。")
    add_heading(doc, "正式资源测量", 2)
    add_body(doc, "资源测量在 test split 上运行，baseline 和 candidate 都使用 CPU FP32 与 4 个 CPU 线程。每个模型预热 1 次，然后在独立进程中正式运行 3 次。时间取三次的中位数，峰值 RSS 取三次最大值。")
    add_table(
        doc,
        ["资源", "Baseline", "Final model", "限制", "占用比例", "状态"],
        [
            ["CPU time median", "4.209 s", "12.103 s", "<= 5x", "57.50%", "通过"],
            ["Peak RSS", "1.873 GiB", "1.948 GiB", "<= 4 GiB", "48.70%", "通过"],
            ["Assets", "4.168 MiB", "40.345 MiB", "<= 64 MiB", "63.04%", "通过"],
        ],
        [1.4, 1.15, 1.25, 1.05, 1.15, 1.2],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 5,
    )
    add_body(doc, "时间比为 12.103 / 4.209 = 2.875x。资源图中显示 57.5%，是因为 2.875x 占允许的 5x 的 57.5%，不是模型只比 baseline 慢 57.5%。Peak RSS 包含 Python、PyTorch、数据和系统库，因此不会简单地与参数量成比例。")
    add_heading(doc, "最终结果可以写成什么", 2)
    add_body(doc, "推荐表述：在 test split 上，最终模型达到 1.6462 BPB，相比 2.1013 BPB 的课程 baseline 相对降低 21.66%。其 CPU FP32 scoring time 为 baseline 的 2.875 倍，峰值 RSS 为 1.948 GiB，推理资产为 40.345 MiB，全部满足课程限制。")
    add_body(doc, "不推荐表述：我们的模型比 baseline 快 2.875 倍。2.875x 表示耗时是 baseline 的 2.875 倍，即更慢但仍在 5x 预算内。也不要写 test BPB 为 1.622441，因为这个数是 validation BPB。")
    # Part 8
    add_heading(doc, "第八部分 图片来源和制作过程", 1)
    add_heading(doc, "图片是否来自论文或网站", 2)
    add_body(doc, "本项目最终使用的两张主要图片都由本地脚本生成，没有从 SCI、CVPR、ACL 论文、网页、搜索引擎或图库复制任何图片元素。学术论文只作为排版习惯参考，例如白底、矢量输出、颜色数量少、面板标号、浅灰网格、图中直接标注和 lower is better 提示。")
    add_table(
        doc,
        ["图片", "生成脚本", "直接数据来源", "外部图片"],
        [
            ["图 1 架构和数据流", "scripts/architecture_diagram.py", "student.py、train_student.py、evaluate.py、ours_a.json、freeze/test JSON", "无"],
            ["图 2 实验汇总", "scripts/plot.py", f"{len(list((CODE / 'runs').glob('*/metrics.json')))} 个 run 的 metrics/curve、grid.tsv、resources JSON", "无"],
        ],
        [1.6, 1.7, 3.2, 0.7],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_heading(doc, "图 1 的制作步骤", 2)
    add_number(doc, "从 configs/ours_a.json 读取 width、depth、heads、FFN、position、tie 和 dropout。")
    add_number(doc, "从 student.py 提取 embedding、RMSNorm、attention、SwiGLU、residual、final norm 和 tied head 的连接关系。")
    add_number(doc, "从 train_student.py 提取 BatchLoader、shifted targets、cross-entropy、AdamW、cosine schedule 和 checkpoint。")
    add_number(doc, "从 evaluate.py 提取 independent windows、log probabilities、target NLL 和 BPB 聚合。")
    add_number(doc, "从 freeze_manifest.json 与 frozen_test_cpu_fp32.json 填入 seed、hash、target 数和最终 test BPB。")
    add_number(doc, "脚本输出 1400 × 920 的 SVG。Word 中插入的是由同一 SVG 高分辨率转换的 PNG，原始 SVG 仍可编辑。")
    add_code(doc, "python scripts/architecture_diagram.py")
    add_heading(doc, "图 2 的制作步骤", 2)
    add_number(doc, "collect.py 扫描所有正式 run 的 metrics.json，生成统一字段的 grid.tsv。")
    add_number(doc, "plot.py 读取 curve.json；旧曲线若缺少 processed_targets，则用 step × batch_size × 256 恢复横轴。")
    add_number(doc, "面板 a 选择能回答 block、容量和训练长度问题的代表曲线。")
    add_number(doc, "面板 b 筛选同一 w320 d8、9.83M-target 预算下的 learning-rate runs。")
    add_number(doc, "面板 c 使用参数量、BPB 和 targets 同时表达 quality-capacity-training trade-off。")
    add_number(doc, "面板 d 读取资源 JSON，将 time ratio、RSS 和 MiB 除以课程上限。")
    add_number(doc, "plot.py 直接输出 1200 × 870 SVG，保证放大后线条和文字不失真。")
    add_code(doc, "python scripts/collect.py\npython scripts/plot.py")
    add_heading(doc, "如何在自己的报告中注明图片来源", 2)
    add_body(doc, "推荐图注：Source: generated by the author from local experiment metrics and resource measurements. 对架构图可以进一步写：The diagram was generated from the submitted model, training and evaluation code. 不需要把 CVPR 或 ACL 论文列为图片来源，因为图形并未复制它们；如果你在正文中讨论 Transformer、RMSNorm 或 SwiGLU 的理论来源，则应引用对应论文。")
    add_heading(doc, "图片可编辑程度", 2)
    add_body(doc, "Word 中的图片本身以高分辨率 PNG 插入，便于稳定显示。真正可编辑的源文件是 code/results 下的 SVG，以及 scripts/plot.py 和 scripts/architecture_diagram.py。你可以修改脚本中的标题、颜色、坐标范围、面板文字和数据筛选，再重新生成 SVG。")
    # Part 9
    add_heading(doc, "第九部分 如何独立撰写最终报告", 1)
    add_heading(doc, "建议的八页结构", 2)
    add_table(
        doc,
        ["页码建议", "内容", "需要回答的问题", "推荐证据"],
        [
            ["1", "标题 摘要 主要结果", "最终做到了什么", "Test BPB 与资源"],
            ["2", "任务 指标 协议", "怎样评分和防止泄漏", "BPB 公式与 split 规则"],
            ["3", "模型和数据流", "模型为什么这样设计", "图 1 和配置表"],
            ["4", "训练与实验设置", "怎样保证公平比较", "targets、seed、optimizer"],
            ["5", "主结果", "block、容量、训练量分别有什么作用", "主结果表和图 2a c"],
            ["6", "超参数和消融", "哪些尝试有效或无效", "LR、dropout、weight decay"],
            ["7", "冻结 test 和资源", "结果是否可信且合规", "freeze hash、test、resources"],
            ["8", "限制 结论 披露 参考文献", "哪些结论不能过度推广", "双 seed、58M 决策、AI disclosure"],
        ],
        [0.8, 1.65, 2.65, 2.1],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_heading(doc, "摘要应该包含什么", 2)
    add_number(doc, "一句说明任务和主要约束。")
    add_number(doc, "一句说明最终模型架构和参数量。")
    add_number(doc, "一句给出 test BPB、baseline 和相对改进。")
    add_number(doc, "一句给出 CPU、内存和资产合规结果。")
    add_number(doc, "一句概括最重要的实验结论，例如容量和训练长度比额外 regularization 更重要。")
    add_body(doc, "摘要不要写实验过程的全部细节，也不要只写 validation 分数。所有数字应能在正文表格或结果 JSON 中找到。")
    add_heading(doc, "方法部分如何写", 2)
    add_body(doc, "先从 baseline 的限制出发：约 1.09M 参数的模型在 9.83M targets 下 validation BPB 为 2.0711。然后说明最终模型如何用 RMSNorm、SwiGLU、无 bias 线性层和 tied weights 提高参数效率，再说明为什么选择 width 320、depth 8：它在同等训练量下明显降低 BPB，同时正式资源测量仍有余量。")
    add_body(doc, "方法部分不能只罗列名词。每个组件至少写出输入输出、基本作用以及为何适合本项目。例如 RMSNorm 的作用是通过均方根缩放稳定 residual 优化；tied weights 的作用是共享输入 embedding 和输出矩阵，从而节省资产。")
    add_heading(doc, "实验部分如何形成论证", 2)
    add_body(doc, "推荐按“问题 受控比较 观察 结论”组织。示例：问题是高效 block 是否有效；受控比较为 s0_lr1e3 与 s_lr1e3；观察是 BPB 从 2.101695 降至 2.040526；结论是 RMSNorm + SwiGLU 组合在相近参数量和相同训练预算下有效。")
    add_table(
        doc,
        ["你想写的结论", "直接证据", "稳妥写法", "应避免的过度结论"],
        [
            ["高效 block 有效", "s0 2.101695 vs s 2.040526", "组合降低约 2.91%", "RMSNorm 单独贡献 2.91%"],
            ["扩大容量有效", "1.082M 2.040526 vs 10.571M 1.820432", "更大模型降低 10.79%", "只有 width 决定收益"],
            ["延长训练有效", "9.83M 1.820432 vs 28.91M 1.622441", "同模型长训练降低 10.88%", "无限延长都会同样改善"],
            ["结果有稳定性", "seed 17 1.622441 vs seed 137 1.631858", "第二个 seed 复现主要趋势", "已经准确估计方差"],
            ["满足资源限制", "2.875x 1.948 GiB 40.345 MiB", "三项都低于上限", "模型比 baseline 更快"],
        ],
        [1.35, 2.0, 2.25, 1.6],
    )
    add_heading(doc, "结果表和图片怎样配合", 2)
    add_body(doc, "表格负责给出精确数字，图片负责帮助读者看趋势。正文先提出问题，再引用表格或图片，最后解释观察。不要让图片独立存在而不在正文讨论，也不要把图中已经清楚显示的每个数字再逐个重复。")
    add_heading(doc, "限制部分应该主动承认什么", 2)
    add_bullet(doc, "v1 最终配方只有两个 seed；补充 RoPE 长训练有三个 seed，方差估计仍然有限。")
    add_bullet(doc, "没有运行 58M targets，因此不能完全排除极长训练的微小收益。")
    add_bullet(doc, "本机环境为 Python 3.13.12，而课程 README 推荐 3.12。")
    add_bullet(doc, "资源测量来自单台机器，绝对秒数依赖硬件，因此同机 baseline 比值更重要。")
    add_bullet(doc, "私有 Git 仓库与 checkpoint release 只解决可追溯发布；课程网站提交仍需人工确认。")
    add_heading(doc, "AI 使用披露应该怎样写", 2)
    add_body(doc, "可以写明使用 Codex 协助诊断代码、安排实验、修复脚本、整理结果、生成图表和编排文档；同时说明全部训练与评测由项目脚本实际执行，模型选择依据为保存的 validation 指标，AI 没有提供外部训练数据、预训练权重或 test 标签。最终提交者仍对代码与结论负责。")
    # Part 10
    add_heading(doc, "第十部分 复现命令 文件位置和检查清单", 1)
    add_heading(doc, "进入环境", 2)
    add_code(doc, "cd MP1_student_starter/code\nsource .venv_project1/bin/activate")
    add_heading(doc, "运行契约测试", 2)
    add_code(doc, "python -m unittest discover -s tests -v")
    add_heading(doc, "重建结果表和两张图", 2)
    add_code(doc, "python scripts/collect.py\npython scripts/plot.py\npython scripts/architecture_diagram.py")
    add_heading(doc, "复算冻结 validation", 2)
    add_code(doc, "python evaluate.py --checkpoint runs/a_e8_lr1e3/checkpoint.pt --device cpu --precision fp32 --threads 4 --split validation --output results/2026-09-23/frozen_validation_cpu_fp32.json")
    add_heading(doc, "复算 test 的注意事项", 2)
    add_body(doc, "课程协议上已经完成冻结后的正式 test，不应为了调整报告而反复运行并据此修改模型。只有在检查软件兼容性或课程明确允许复核时才重新运行；重新运行也不能改变冻结 checkpoint。")
    add_heading(doc, "正式资源测量命令", 2)
    add_code(doc, "python scripts/measure.py runs/a_e8_lr1e3/checkpoint.pt --baseline runs/R0_baseline/checkpoint.pt --threads 4 --split test --warmups 1 --repeats 3 --output results/2026-09-23/resources_test.json")
    add_heading(doc, "写报告前核对清单", 2)
    checklist = [
        "所有 test 数字来自 frozen_test_cpu_fp32.json 或 resources_test.json。",
        "Validation BPB 1.622441 与 test BPB 1.646152 没有混淆。",
        "时间 2.875x 被正确解释为耗时比，而不是速度提升。",
        "图注说明图片由本地实验数据和脚本生成。",
        "每个主结论都有受控对照或明确限制。",
        "负面实验 dropout 和 weight decay 被保留，而不是只展示成功结果。",
        "报告说明 test 在模型冻结后才运行。",
        "资源表包含 baseline、candidate、课程上限和通过状态。",
        "AI assistance disclosure 已写入。",
        "Git release 与网站提交若仍未完成，不应写成已完成。",
    ]
    for item in checklist:
        add_bullet(doc, "□ " + item)
    add_heading(doc, "关键文件索引", 2)
    add_table(
        doc,
        ["文件", "内容"],
        [
            ["code/results/grid.tsv", f"{len(list((CODE / 'runs').glob('*/metrics.json')))} 个正式实验的统一汇总"],
            ["code/results/2026-09-23/seed_comparison.json", "两个最终 seed 的选择证据"],
            ["code/results/2026-09-23/freeze_manifest.json", "test 前冻结配置和哈希"],
            ["code/results/2026-09-23/frozen_test_cpu_fp32.json", "最终 test BPB 和 NLL"],
            ["code/results/2026-09-23/resources_test.json", "正式 test 资源测量"],
            ["code/results/figures/experiment_summary.svg", "可编辑实验汇总图"],
            ["code/results/2026-09-23/architecture_data_flow.svg", "可编辑架构与数据流图"],
            ["output/pdf/MP1_Final_Report.pdf", "10 页正式报告参考"],
        ],
        [3.95, 3.25],
    )
    # Appendix
    add_heading(doc, "附录 指标词汇 哈希和常见问题", 1)
    add_heading(doc, "指标词汇", 2)
    add_table(
        doc,
        ["术语", "含义", "本项目中的用法"],
        [
            ["BPB", "每个原始字节平均需要的 bit 数", "主质量指标，越低越好"],
            ["Token PPL", "以 token 为单位的 perplexity", "补充指标，不用于主要排名"],
            ["NLL", "正确 token 的负对数概率之和", "BPB 的分子来源"],
            ["Processed targets", "训练中参与 next-token loss 的目标数", "公平比较训练预算"],
            ["Checkpoint", "训练完成后保存的模型权重和配置", "最终选中 seed 17 权重"],
            ["Asset MiB", "推理所需未压缩资产大小", "课程限制为 64 MiB"],
            ["Peak RSS", "进程最大 resident memory", "课程限制为 4 GiB"],
            ["Freeze", "test 前锁定模型、代码和配置", "防止 test 参与选择"],
        ],
        [1.35, 3.2, 2.65],
    )
    add_heading(doc, "常见问题", 2)
    faqs = [
        ("为什么 validation 比 test 更低", "两个 split 的文本分布和难度不完全相同。只要模型选择没有使用 test，这种差异是正常的。"),
        ("为什么最终 test 只运行一次", "防止观察 test 后继续调参。资源测量可以重复计时，但使用同一冻结模型，不改变配置。"),
        ("为什么训练用 MPS BF16 而最终用 CPU FP32", "MPS BF16 提高训练速度；课程最终评分要求 CPU FP32，因此最终分数和资源都在 CPU FP32 下测量。"),
        ("为什么 baseline RSS 和大模型接近", "进程 RSS 包含运行时、框架、数据和内存映射，不只包含参数。资产 MiB 更直接反映模型权重大小。"),
        ("为什么不选择更大的 13M 或 15M 模型", "相同 9.83M-target 预算下，w320 d10 为 1.872762，w384 d8 为 1.829156，都没有超过 w320 d8 的 1.820432；后者资产还接近 64 MiB 上限。"),
        ("图片能否直接放入报告", "可以，但应保留图注和来源说明，并在正文解释每个面板回答的问题。"),
        ("Word 中的图是否完全可编辑", "插入的是高分辨率 PNG。图形源文件 SVG 和生成脚本完全可编辑，修改后重新插入即可。"),
        ("自己写报告最容易犯什么错误", "把 validation 当 test、把 2.875x 写成速度提升、只展示成功实验、把组合消融归因到单个组件、没有说明冻结顺序。"),
    ]
    for q, a in faqs:
        para = doc.add_paragraph()
        add_text(para, "问 " + q, bold=True)
        ans = doc.add_paragraph(style="Body Text")
        add_text(ans, "答 " + a)

    add_heading(doc, "参考文献", 2)
    refs = [
        "Merity S et al. Pointer Sentinel Mixture Models. ICLR 2017. WikiText benchmark.",
        "Vaswani A et al. Attention Is All You Need. NeurIPS 2017.",
        "Zhang B and Sennrich R. Root Mean Square Layer Normalization. NeurIPS 2019.",
        "Shazeer N. GLU Variants Improve Transformer. arXiv 2002.05202. 2020.",
        "Su J et al. RoFormer. arXiv 2104.09864. 2021.",
        "Loshchilov I and Hutter F. Decoupled Weight Decay Regularization. ICLR 2019.",
    ]
    for ref in refs:
        add_bullet(doc, ref)

    add_heading(doc, "最终状态", 2)
    add_body(doc, "冻结 v1、最终 test、正式资源测量和补充 validation-only 消融均已完成。私有 Git 仓库保留逻辑提交历史，checkpoint 通过私有 release 发布而不进入 Git 历史。课程网站提交需要由提交者在截止时间前人工完成。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    build_document()
