#!/usr/bin/env python3
"""Build the complete editable Chinese DOCX submission report."""

from __future__ import annotations

import shutil
import sys
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import word_report_common as base


SOURCE = Path("/Users/echomisty/Desktop/DASE7506 Project 1 TECH Report.docx")
PROJECT = Path(__file__).resolve().parents[2]
OUTPUT = PROJECT / "output" / "docx" / "DASE7506_Project_1_TECH_Report_中文修订版.docx"
ARCHITECTURE_PNG = PROJECT / "tmp" / "docx" / "assets" / "architecture_data_flow.png"
FIGURE_DIR = PROJECT / "tmp" / "tech_report_figures"
OPTIMIZATION_PNG = PROJECT / "tmp" / "experiment_decision_flow" / "flow-1.png"
ABLATION_PNG = FIGURE_DIR / "learning_rate.png"
SUPPLEMENT_PNG = FIGURE_DIR / "rope_three_seeds.png"

BODY_SIZE = 10.3
TABLE_SIZE = 8.4
CAPTION_SIZE = 8.3
NAVY = "1F4E78"
LIGHT_BLUE = "F3F7FB"
BORDER = "D9D9D9"


def clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_style_font(style, *, latin: str, east_asia: str, size: float, bold: bool = False,
                   color: str = "000000") -> None:
    base.set_style_font(style, latin=latin, east_asia=east_asia, size=size, bold=bold, color=color)


def format_body(paragraph, *, first_line: bool = True, after: float = 2.6) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(BODY_SIZE * 2 if first_line else 0)
    fmt.line_spacing = 1.15
    fmt.space_after = Pt(5)
    fmt.widow_control = True


def add_body(document: Document, text: str, *, first_line: bool = True, bold_lead: str | None = None):
    paragraph = document.add_paragraph(style="Normal")
    format_body(paragraph, first_line=first_line)
    if bold_lead and text.startswith(bold_lead):
        run = paragraph.add_run(bold_lead)
        base.set_run_font(run, size=BODY_SIZE, bold=True)
        run = paragraph.add_run(text[len(bold_lead):])
        base.set_run_font(run, size=BODY_SIZE)
    else:
        run = paragraph.add_run(text)
        base.set_run_font(run, size=BODY_SIZE)
    return paragraph


def add_heading(document: Document, text: str, *, page_break: bool = False):
    paragraph = document.add_paragraph(text, style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.page_break_before = page_break
    return paragraph


def add_subheading(document: Document, text: str):
    paragraph = document.add_paragraph(text, style="Heading 2")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    return paragraph


def add_caption(document: Document, text: str, *, before: bool = True):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2 if before else 1)
    paragraph.paragraph_format.space_after = Pt(2 if before else 4)
    paragraph.paragraph_format.keep_with_next = before
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    base.set_run_font(run, size=CAPTION_SIZE, bold=True, color="333333")
    return paragraph


def fill_cell(cell, text: str, *, header: bool = False, align: str = "left") -> None:
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    base.set_cell_margins(cell, top=58, start=78, bottom=58, end=78)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    run = paragraph.add_run(text)
    base.set_run_font(run, size=TABLE_SIZE, bold=header, color="FFFFFF" if header else "222222")
    if header:
        base.shade_cell(cell, NAVY)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float],
              aligns: list[str] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    base.set_table_width(table, sum(widths))
    base.set_table_borders(table)
    aligns = aligns or ["left"] * len(headers)
    for col, text in enumerate(headers):
        table.columns[col].width = Inches(widths[col])
        base.set_cell_width(table.rows[0].cells[col], widths[col])
        fill_cell(table.rows[0].cells[col], text, header=True, align="center")
    base.set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        base.prevent_row_split(row)
        for col, text in enumerate(values):
            base.set_cell_width(row.cells[col], widths[col])
            fill_cell(row.cells[col], text, align=aligns[col])
            if row_index % 2 == 1:
                base.shade_cell(row.cells[col], LIGHT_BLUE)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.5
    return table


def set_picture_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_figure(document: Document, image_path: Path, caption: str, *, width: float = 5.62,
               alt_title: str, alt_description: str):
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_with_next = True
    inline_shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    set_picture_alt_text(inline_shape, alt_title, alt_description)
    add_caption(document, caption, before=False)
    return inline_shape


def restore_preserved_package_parts(source: Path, output: Path) -> None:
    """Keep opaque source parts while retaining the new body, styles and images."""
    editable = {
        "[Content_Types].xml",
        "word/document.xml",
        "word/styles.xml",
        "word/_rels/document.xml.rels",
    }
    temp_output = output.with_suffix(".preserve-parts.docx")
    with ZipFile(source, "r") as source_zip, ZipFile(output, "r") as final_zip, ZipFile(temp_output, "w") as out_zip:
        source_names = set(source_zip.namelist())
        for info in final_zip.infolist():
            if info.filename in source_names and info.filename not in editable:
                data = source_zip.read(info.filename)
            else:
                data = final_zip.read(info.filename)
            out_zip.writestr(info, data)
    temp_output.replace(output)


def build() -> None:
    for required in (SOURCE, ARCHITECTURE_PNG, OPTIMIZATION_PNG, ABLATION_PNG, SUPPLEMENT_PNG):
        if not required.exists():
            raise FileNotFoundError(required)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    clear_document_body(document)
    for section in document.sections:
        grid = section._sectPr.find(qn("w:docGrid"))
        if grid is not None:
            section._sectPr.remove(grid)
    summary = json.loads((PROJECT / "code/results/2026-09-25/supplementary_summary.json").read_text())
    long_runs = summary["rope_long_training"]
    resources = summary["validation_resource_measurement"]["models"][0]
    saved_runs = [json.loads(p.read_text()) for p in (PROJECT / "code/runs").glob("*/metrics.json") if "smoke" not in str(p).lower()]
    train_seconds = sum(r.get("train_seconds", 0) for r in saved_runs)

    set_style_font(document.styles["Normal"], latin=base.BODY_LATIN, east_asia=base.BODY_CN,
                   size=BODY_SIZE, color="000000")
    set_style_font(document.styles["Title"], latin=base.TITLE_LATIN, east_asia=base.HEADING_FONT,
                   size=18.5, color="000000")
    set_style_font(document.styles["Heading 1"], latin=base.HEADING_FONT, east_asia=base.HEADING_FONT,
                   size=14.5, bold=True, color="000000")
    set_style_font(document.styles["Heading 2"], latin=base.HEADING_FONT, east_asia=base.HEADING_FONT,
                   size=10.8, bold=True, color="000000")
    document.styles["Heading 1"].paragraph_format.space_before = Pt(7)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(3)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(5)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(2)

    title = document.add_paragraph("DASE7506 Project 1 语言模型 BPB 技术报告", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(7)
    run = subtitle.add_run("WikiText 2 模型优化与三 seed 验证")
    base.set_run_font(run, size=10.5, bold=False, color="000000")

    add_body(
        document,
        "摘要  本项目从随机初始化训练 decoder-only Transformer。通过 block 对照、容量筛选和训练长度实验，冻结 v1 的 test BPB 降至 1.646152，较课程基线降低 21.66%，三项资源指标均满足上限。补充实验支持继续考察 RoPE；三 seed 的 CPU FP32 validation 均值为 1.600116，候选资源检查已完成，其 test 尚未评估。",
        first_line=False,
        bold_lead="摘要  ",
    )

    add_heading(document, "1 背景与任务目标")
    # Preserve the two substantive opening sentences from the user's source file.
    add_body(
        document,
        "从零开始训练语言模型，改进其架构或训练过程，并在资源限制内实现最低的可复现测试“每字节比特数”（bits per byte, BPB）。本项目在课程固定的 WikiText-2、BPE-2048 tokenizer、256-token 上下文和因果评测协议下，从随机初始化训练一个紧凑 decoder-only Transformer。",
    )
    add_body(
        document,
        "目标是在不使用外部文本、预训练权重或 test 调参的前提下，最小化完整 test split 的 bits per byte (BPB)。对每个独立窗口，模型根据前文预测下一个 token；除 split 第一个 token 外，每个 target 恰好评分一次。",
    )
    add_caption(document, "表 1 作业的固定条件与硬性上限")
    add_table(
        document,
        ["类别", "要求", "本项目处理"],
        [
            ["数据", "只用 train 学习；validation 选择；test 不调参", "固定 split 用途，test 前冻结"],
            ["评测", "BPE-2048；context 256；独立 causal 窗口；CPU FP32", "固定 tokenizer 与 evaluator"],
            ["CPU 时间", "不超过 baseline 的 5 倍", "最终 2.875 倍"],
            ["内存", "峰值 RSS 不超过 4 GiB", "最终 1.948 GiB"],
            ["推理资产", "未压缩资产不超过 64 MiB", "最终 40.345 MiB"],
            ["提交", "报告不超过 10 页；代码与 checkpoint 可复现；披露 AI 帮助", "正文、复现证据和披露均覆盖"],
        ],
        [0.90, 2.62, 2.25],
        ["center", "left", "left"],
    )

    add_heading(document, "2 评价指标与选择协议", page_break=True)
    add_body(
        document,
        "设 B 为 split 原始 UTF-8 文本的字节数。固定 evaluator 对所有被评分 target 累加负对数似然，再把 nats 转换为 bits 并除以 B。",
    )
    base.add_bpb_equation(document)
    add_body(
        document,
        "BPB 越低越好。完整 split 的总负对数似然必须先汇总再除以字节数，不能对 batch 或窗口的 BPB 做简单平均。Token perplexity 只作为辅助指标，因为课程排名以 BPB 为准。",
    )
    add_caption(document, "表 2 报告指标与判断方式")
    add_table(
        document,
        ["指标", "计算或测量", "用途"],
        [
            ["BPB", "总 NLL ÷ ln(2) ÷ UTF-8 bytes", "主要质量指标，越低越好"],
            ["Token PPL", "exp（总 NLL ÷ target 数）", "固定 tokenizer 下的辅助解释"],
            ["CPU 时间比", "同机候选时间 ÷ baseline 时间", "判断是否超过 5 倍"],
            ["峰值 RSS", "独立评测进程的最大常驻内存", "判断是否超过 4 GiB"],
            ["推理资产", "推理所需文件的未压缩大小", "判断是否超过 64 MiB"],
        ],
        [1.12, 2.75, 1.90],
        ["center", "left", "left"],
    )
    add_subheading(document, "2.1 数据用途和冻结")
    add_body(
        document,
        "train 只用于更新参数。validation 用于选择结构、学习率、正则、seed 和 checkpoint。选定 seed 17 后，项目先记录 checkpoint、student.py、evaluate.py 和 tokenizer.json 的 SHA-256，再运行 student-model test。测试结果没有回流到 v1 的结构或超参数选择。",
    )
    add_body(
        document,
        "公平比较时固定数据、tokenizer、context、batch size 和 evaluator。短实验统一处理 9,830,400 个训练 targets。模型规模或训练时长发生变化时，报告同时给出 BPB 与计算成本，避免把额外计算误写成单一机制的收益。",
    )
    add_caption(document, "表 3 主线实验的规模与训练预算")
    add_table(document, ["实验", "参数量", "训练 targets", "Validation BPB", "结论"], [
        ["课程 baseline", "1.088M", "9.83M", "2.071087", "复现基准"],
        ["小型 LN + GELU", "1.084M", "9.83M", "2.101695", "block 对照"],
        ["小型 RMS + SwiGLU", "1.082M", "9.83M", "2.040526", "较小型对照低 2.91%"],
        ["w320 d8", "10.571M", "9.83M", "1.820432", "容量带来 10.79% 改进"],
        ["w320 d8 长训练", "10.571M", "28.91M", "1.622441", "再改善 10.88%"],
    ], [1.40, 0.82, 0.92, 1.00, 1.63], ["left", "center", "center", "center", "left"])

    add_heading(document, "3 模型与训练方法", page_break=True)
    add_body(
        document,
        "冻结的 v1 模型包含 8 个 decoder block，宽度为 320，使用 4 个 attention heads。每个 block 采用 pre-RMSNorm、causal scaled dot-product attention 和 SwiGLU FFN，线性层不使用 bias。输入 embedding 与输出 head 共享权重，位置表示使用 learned positional embeddings，dropout 为 0。模型共有 10,570,560 个参数。",
    )
    add_figure(
        document,
        ARCHITECTURE_PNG,
        "图 1 冻结 v1 的模型结构、训练数据流与评测流程",
        width=5.60,
        alt_title="冻结 v1 架构和数据流",
        alt_description="从 WikiText-2 训练文本、固定 tokenizer、八层 Transformer、AdamW 更新到冻结评测和 BPB 汇总的流程图。",
    )
    add_body(
        document,
        "训练样本由 257 个连续 tokens 构成，前 256 个作为输入，后 256 个作为目标。batch size 为 32，每步处理 8,192 个 targets。优化器为 AdamW，peak learning rate 为 1e-3，weight decay 为 0.1，betas 为 (0.9, 0.95)。学习率先 warm up 3%，再按 cosine schedule 降到峰值的 10%，gradient clipping 为 1.0。训练使用 Apple MPS BF16，正式分数使用 CPU FP32 重算。",
    )
    add_body(document, "各配方和各 seed 均独立从随机初始化训练。28.91M-target 长实验重新训练并延长学习率 schedule，并非在 9.83M checkpoint 上直接续训。冻结 v1 采用末尾权重；后续 RoPE 实验另行加入预声明检查点的最佳权重保存。")

    add_heading(document, "4 BPB 优化过程与判断方法", page_break=True)
    add_body(
        document,
        "图 2 将冻结 v1 的开发与后续 RoPE 候选串成同一条决策路径。箭头表示基于上一轮结果安排下一轮实验，不表示直接续训；各组从随机初始化训练。",
    )
    add_figure(
        document,
        OPTIMIZATION_PNG,
        "图 2 实验数据、结果含义与后续决策的完整流程",
        width=5.62,
        alt_title="BPB 优化路径",
        alt_description="八步完整实验决策图，合并冻结 v1 前的 block、容量、超参数、训练长度和 seed 选择，以及冻结后的消融、RoPE、最佳 checkpoint 与三 seed 复核。每一步显示实验数据、含义和后续决策。",
    )
    add_body(
        document,
        "课程 baseline 的 2.071087 是复现起点，不能替代图中 2.101695 的受控比较。组合 block 先取得 2.91% 改进，随后扩大 width 和 depth 再取得 10.79% 改进，因此预算转向 w320 d8 的超参数与训练长度。独立组件的贡献由第 7 节补充消融检验。",
    )

    add_heading(document, "5 超参数与训练长度的决策依据", page_break=True)
    add_body(
        document,
        "在确定 w320 d8 后，我们先检查学习率是否限制优化。六个短实验均使用 seed 17 和 9.83M targets；1e-3 得到最低 BPB，而 7.5e-4 仅差 0.006146。1.25e-3、2e-3 和 4e-3 均更差，因此固定 1e-3 进入长训练，不再扩大高学习率搜索。",
    )
    add_figure(
        document,
        ABLATION_PNG,
        "图 3 固定训练预算下的学习率扫描",
        width=5.62,
        alt_title="学习率训练轨迹与最终 BPB",
        alt_description="六个学习率的实际完整 validation BPB，横轴为对数坐标，1e-3 的观测值最低。连线仅帮助读取趋势。",
    )
    add_caption(document, "表 4 不同结果如何改变后续方向")
    add_table(
        document,
        ["比较", "结果", "判断", "后续动作"],
        [
            ["Peak LR", "1e-3 为 1.820432", "扫描最低点", "固定 1e-3"],
            ["Dropout 0.05", "+0.033451 BPB", "明显变差", "v1 不使用 dropout"],
            ["Weight decay 0.2", "+0.000085 BPB", "未见有意义收益", "保留已有的 0.1"],
            ["Seed 137", "+0.009417 BPB", "两者都支持长训练收益", "按预先规则选更低的 seed 17"],
            ["58M targets", "未运行", "最后约 7.2M 仅改善约 0.0032", "不把训练成本翻倍"],
        ],
        [1.15, 1.22, 1.56, 1.84],
        ["left", "center", "left", "left"],
    )
    add_body(document, "Dropout 0.05 使 BPB 从 1.820432 升至 1.853883，而 weight decay 0.2 几乎没有收益。这说明当前短训练预算下，增强正则化不是优先方向；仅凭这组结果还不能断言模型一定欠拟合。我们保留 dropout 0 和 weight decay 0.1，转而测试更长训练。")
    add_body(document, "将完整训练预算从 9.83M 扩到 28.91M 后，BPB 从 1.820432 降至 1.622441，说明这一配方仍能从更多更新中受益。末尾约 7.2M targets 仅从 1.625607 改善到 1.622441，收益已明显减缓，因此没有继续运行成本近乎翻倍的 58M 配方。")
    add_body(document, "随后加入 seed 137，得到 1.631858。两个 seed 均支持长训练收益，但 0.009417 的差距说明随机初始化仍影响结果。我们按最低 validation 选择 seed 17；这提供跨 seed 支持，尚不足以精确估计总体方差。")

    add_heading(document, "6 冻结后的最终结果与资源", page_break=True)
    add_body(
        document,
        "seed 17 在完整 validation 上优于 seed 137，项目按照预先确定的最低 validation BPB 规则选择该 checkpoint。冻结清单记录哈希后，student-model test 只运行一次。",
    )
    add_caption(document, "表 5 完整 test split 的结果")
    add_table(
        document,
        ["模型", "Test BPB", "Token PPL", "NLL nats", "Targets", "UTF-8 bytes"],
        [
            ["课程 baseline", "2.101265", "80.8479", "1,881,798.91", "428,405", "1,292,013"],
            ["冻结 v1", "1.646152", "31.2239", "1,474,220.05", "428,405", "1,292,013"],
        ],
        [1.08, 0.78, 0.78, 1.13, 0.92, 1.08],
        ["left", "center", "center", "right", "right", "right"],
    )
    add_body(
        document,
        "冻结 v1 相对 baseline 绝对降低 0.455113 BPB，相对降低 21.66%。两个模型评分的 targets 和原始字节数相同，因此差异来自概率质量，而不是评分覆盖变化。",
    )
    add_caption(document, "表 6 test 资源测量")
    add_table(
        document,
        ["指标", "Baseline", "冻结 v1", "课程上限", "结论"],
        [
            ["CPU FP32 时间中位数", "4.209 s", "12.103 s", "不超过 5×", "2.875×，通过"],
            ["正式测量范围", "4.209-4.238 s", "12.096-12.164 s", "同一流程", "预热 1 次，重复 3 次"],
            ["峰值 RSS", "1.873 GiB", "1.948 GiB", "4 GiB", "通过"],
            ["推理资产", "4.168 MiB", "40.345 MiB", "64 MiB", "通过"],
        ],
        [1.47, 1.03, 1.08, 1.02, 1.17],
        ["left", "center", "center", "center", "left"],
    )
    add_subheading(document, "6.1 训练和搜索成本")
    add_body(
        document,
        f"最终 v1 的 seed 17 训练用时为 1,143 s，seed 137 为 1,093 s。当前保留 metrics 的 {len(saved_runs)} 个非 smoke runs 合计记录 {train_seconds:,.0f} s，约 {train_seconds/3600:.2f} 小时。这一统计包含冻结后的补充实验，按各 run 的 train_seconds 字段汇总；不包含单独的 CPU 复核、资源测量、数据准备和文档生成。",
    )
    add_body(document, "RoPE 三次长训练分别记录约 1,436、1,399 和 1,349 s。虽然最终权重来自 19.66M targets，三次搜索实际都运行到 28.91M，因此披露成本时使用完整搜索量。训练时间来自单机运行记录，适合说明本项目投入，不能直接视为不同机制的硬件无关成本。")

    add_heading(document, "7 补充消融如何改变后续方向", page_break=True)
    add_body(
        document,
        "以下实验在 v1 test 之后完成，只用 validation 解释机制并筛选后续候选，对应图 2 下半部分。短实验固定 9.83M targets，数值来自 MPS 记录。额外收益不足 0.01 BPB 时不优先投入长训练；这一经验阈值用于控制搜索成本，不等同于统计显著性检验。",
    )
    add_caption(document, "表 7 相同预算下的 Normalization 与 FFN 消融")
    add_table(document, ["Normalization", "FFN", "Validation BPB", "相对 LN＋GELU"], [
        ["LayerNorm", "GELU", "2.101695", "0"],
        ["RMSNorm", "GELU", "2.095283", "−0.006412"],
        ["LayerNorm", "SwiGLU", "2.045564", "−0.056130"],
        ["RMSNorm", "SwiGLU", "2.040526", "−0.061169"],
    ], [1.50,1.02,1.50,1.75], ["left","center","center","center"])
    add_body(
        document,
        "固定 LayerNorm 时，SwiGLU 改善 0.056130 BPB；固定 GELU 时，RMSNorm 只改善 0.006412。因此 block 组合的主要收益由 SwiGLU 支持，RMSNorm 的独立收益仍较小。该组使用 seed 17，不能把单次差值写成已证实的跨 seed 稳定效应。",
    )
    add_caption(document, "表 8 补充实验的观察和改进决策")
    add_table(document, ["问题", "实验观察", "判断与后续动作"], [
        ["收益是否来自训练器", "基准 2.072376；无放回 2.071598；再改 beta2 为 2.103114；完整配方 2.098476", "采样近乎中性，beta2 在小 baseline 上变差。训练器并未单独解释结构收益，继续检查 block。"],
        ["容量是否还能增加", "w320 d8 为 1.820432；d10 为 1.872762；w384 d8 为 1.829156", "更大模型在相同预算下未胜出；后者资产 57.422 MiB。保留 w320 d8，把预算转向表示层。"],
        ["位置与权重共享", "seed 17：learned tied 1.820432，untied 1.786678，RoPE tied 1.754858；RoPE 双 seed 均值 1.748021", "RoPE 优先进入长训练。其后 untied 的配对平均额外收益约 0.00675，低于阈值，仍使用 tied。"],
        ["EMA 是否值得保留", "同一次训练的 live 1.735489，EMA 1.734143，改善 0.001346", "仅比较同一优化轨迹，避免把不同 run 波动算作 EMA 收益。改善过小，未推进长训练。"],
        ["MTP 是否值得保留", "MTP 为 1.739964；对照 RoPE untied seed 17 为 1.745039，改善 0.005075", "改善不足 0.01，未推进。单次运行时间不足以确认固定成本增幅，需要受控复测。"],
    ], [1.12,2.23,2.42], ["left","left","left"])

    add_heading(document, "8 RoPE 三 seed 与最佳 checkpoint", page_break=True)
    add_body(
        document,
        "第一次 RoPE 长训练在 19.66M targets 达到 1.620751，末尾却回升至 1.634776，当时只保存了最后权重。这个观察直接改变了训练流程：后续重跑前声明 9.83M、19.66M、28.91M 三个验证点，并保存其中最低 validation BPB 的权重。",
    )
    add_figure(document, SUPPLEMENT_PNG, "图 4 RoPE 三个 seed 的验证曲线与选择位置", width=5.10,
               alt_title="RoPE 三 seed 验证曲线", alt_description="三个独立 RoPE 长训练的真实 MPS 验证点，横轴为累计训练 targets；三者都在 19.66M 处优于末尾。完整 CPU FP32 复核结果在下表列出。")
    add_caption(document, "表 9 最佳权重的 CPU FP32 复核与训练末尾比较")
    add_table(document, ["Seed", "Search targets", "Selected targets", "CPU FP32 BPB", "末尾 MPS BPB"], [
        [str(r["seed"]), "28.91M", "19.66M", f'{r["validation_bpb_cpu_fp32"]:.6f}', f'{r["final_validation_bpb_mps_bf16"]:.6f}'] for r in long_runs["runs"]
    ], [0.52,1.10,1.15,1.45,1.55], ["center"]*5)
    add_body(
        document,
        f"三 seed 的 CPU FP32 validation 均值为 {long_runs['mean_bpb']:.6f}，样本标准差为 {long_runs['sample_stdev_bpb']:.6f}。三次均选中 step 2400 的权重，但搜索均持续至 28.91M。验证曲线支持保存最佳 checkpoint，而不是固定采用末尾权重。",
    )
    add_body(document, "同 seeds 17、137 的 RoPE 均值为 1.594506，v1 为 1.627150，改善 0.032644。这是两 seed 子集的整体方案比较；RoPE 还采用最佳 checkpoint，不能把全部差值归因于位置编码。")
    add_caption(document, "表 10 seed 137 候选的 validation 资源检查")
    add_table(document, ["指标", "实测", "课程上限", "结果"], [
        ["CPU FP32 时间比", f"{resources['time_vs_baseline']:.3f}× baseline", "5×", "通过"],
        ["峰值 RSS", f"{resources['peak_rss_gib_max']:.3f} GiB", "4 GiB", "通过"],
        ["未压缩推理资产", f"{resources['asset_mib']:.3f} MiB", "64 MiB", "通过"],
    ], [1.72,1.75,1.20,1.10], ["left","center","center","center"])
    add_body(document, "seed 137 按最低 CPU FP32 validation 选为候选，BPB 为 1.588095。资源测量使用 validation split、4 线程、1 次预热和 3 次独立进程重复；候选清单已记录 checkpoint 哈希。候选尚未完成独立正式冻结和 test，因此正式成绩仍为 v1 的 1.646152。")

    add_heading(document, "9 正确性与提交要求", page_break=True)
    add_body(
        document,
        "课程固定代码、数据和 tokenizer 通过 manifest 哈希检查。实现不缓存 validation/test 答案，不访问未来 token，不跨窗口保留状态，评测无需联网。契约测试覆盖因果性、概率归一化、样本独立、窗口重置、错位目标梯度及末尾短窗口；训练器测试检查采样、学习率与最佳 checkpoint 保存。RoPE 广播维和曲线横轴修复已记录在实验文档。",
    )
    add_subheading(document, "9.1 复现材料")
    add_body(
        document,
        "冻结 checkpoint 位于 code/runs/a_e8_lr1e3/checkpoint.pt，SHA-256 为 ab9ba8b648513530c528ce00345ba1dd8749351c8cf5516dc5ace84f40ad6657。结果、资源测量、seed 比较和 freeze manifest 保存在 code/results/2026-09-23/。环境为 Python 3.13.12、PyTorch 2.7.1、NumPy 2.5.3 和 tokenizers 0.21.4；发布环境仍建议按课程 README 使用 Python 3.12 复核。",
    )

    add_subheading(document, "9.2 作业要求覆盖")
    add_caption(document, "表 11 课程要求与证据位置")
    add_table(
        document,
        ["课程要求", "本报告或仓库中的证据", "状态"],
        [
            ["从随机初始化训练，只用课程 train", "第 1、3 节；README 数据规则", "完成"],
            ["固定 BPE-2048、context 256 与 evaluator", "第 1、2 节；哈希检查", "完成"],
            ["初始 baseline", "表 3 与表 5", "完成"],
            ["相同 processed targets 的比较", "表 3 的 9.83M 对照", "完成"],
            ["关键机制消融", "表 7 的 2×2 对照与表 8 的筛选", "完成"],
            ["validation 开发、test 前冻结", "第 2、6 节与 freeze manifest", "完成"],
            ["质量与计算代价分析", "表 3、表 6 与停止 58M 的理由", "完成"],
            ["训练成本、seed 与 checkpoint ancestry", "第 3、5、6、8 节；grid.tsv", "完成"],
            ["5× CPU、4 GiB、64 MiB 限制", "表 6", "通过"],
            ["不超过 10 页的报告与 AI 披露", "本报告第 10 节", "完成"],
            ["不可变代码链接与 checkpoint 下载链接", "需在截止前写入最后一次课程提交", "待外部提交"],
            ["课程网站登记与后续 peer review", "29 日前登记；公开后复现另一份提交", "待外部提交"],
        ],
        [2.34, 2.72, 0.71],
        ["left", "left", "center"],
    )
    add_body(
        document,
        "9 月 29 日前须登记学号与 full-test BPB 并完成生成的 GitHub issue；9 月 30 日最终提交须附不可变代码和匹配 checkpoint 下载链接，提供精确安装、训练与评测命令。公开后 7 天内按课程流程复现另一份提交并报告分数。仓库 README 已披露 AI 帮助，网站提交及外部链接状态需在提交时核对。",
    )

    add_heading(document, "10 结论与局限及 AI 披露", page_break=True)
    add_subheading(document, "10.1 结论")
    add_body(
        document,
        "冻结 v1 把 test BPB 从 2.101265 降到 1.646152，并通过时间、内存和资产限制。实验先支持保留高效 block 和扩大容量，再通过学习率、正则与训练曲线确定后续投入。补充对照把主要 block 收益定位到 SwiGLU；位置实验和长训练回退又推动了 RoPE 与最佳 checkpoint 保存的结合。",
    )
    add_body(document, "RoPE 三 seed 的 CPU FP32 validation 均值为 1.600116，最低的 seed 137 为 1.588095，候选资源门槛通过。它说明新方案值得保留，但尚无新 test 结论。负面实验同样改变了路径：更深更宽模型未进入后续训练，untied、EMA 和 MTP 因收益不足而停止。")
    add_subheading(document, "10.2 局限")
    add_body(
        document,
        "v1 只比较两个 seeds，RoPE 长训练比较三个，其他多数短消融只有 seed 17。均值和小幅差异仍有不确定性，0.01 BPB 仅是搜索决策阈值。RoPE 长方案还改变了 checkpoint 选择方式，长训练收益不能全部归因于位置编码。没有运行 58M 实验，也不能排除更长训练经重新调参后仍有收益。",
    )
    add_body(document, "训练在 MPS BF16 上进行，正式分数以 CPU FP32 为准；短筛选不全是 CPU 复核结果。资源来自单机测量，绝对时间会受硬件与系统负载影响；候选资源在 validation 上测量，不能与 v1 test 的绝对秒数直接相减解释开销。Python 3.12 的独立安装复现仍应在发布环境中完成。")
    add_subheading(document, "10.3 AI assistance disclosure")
    add_body(
        document,
        "OpenAI Codex 协助了代码诊断、实验计划、实现复查、脚本修复、结果整理、图表和文档排版。所有训练与评测数值均由本项目脚本产生，并保留 machine-readable evidence。Codex 未提供外部训练文本、预训练权重、validation 答案或 test 答案。提交者负责理解实现、核对证据并完成最终提交。",
    )
    add_subheading(document, "参考资料")
    add_body(
        document,
        "[1] DASE7506 Mini Project 1 Guide，GUIDE.md。\n"
        "[2] DASE7506 MP1 Starter Code README，code/README.md。\n"
        "[3] Vaswani et al. Attention Is All You Need. NeurIPS 2017。\n"
        "[4] Zhang and Sennrich. Root Mean Square Layer Normalization. NeurIPS 2019。\n"
        "[5] Shazeer. GLU Variants Improve Transformer. arXiv:2002.05202, 2020。\n"
        "[6] Su et al. RoFormer Enhanced Transformer with Rotary Position Embedding. arXiv:2104.09864, 2021。\n"
        "[7] Loshchilov and Hutter. Decoupled Weight Decay Regularization. ICLR 2019。",
        first_line=False,
    )

    # Reapply fonts to visible runs created by paragraph.text and table helpers.
    for paragraph in document.paragraphs:
        ppr = paragraph._p.get_or_add_pPr()
        snap = ppr.find(qn("w:snapToGrid"))
        if snap is None:
            snap = OxmlElement("w:snapToGrid")
            ppr.append(snap)
        snap.set(qn("w:val"), "0")
        style_name = paragraph.style.name if paragraph.style else "Normal"
        for run in paragraph.runs:
            if style_name == "Title":
                base.set_run_font(run, latin=base.TITLE_LATIN, east_asia=base.HEADING_FONT)
            elif style_name in {"Heading 1", "Heading 2"}:
                base.set_run_font(run, latin=base.HEADING_FONT, east_asia=base.HEADING_FONT)
            else:
                current_color = run.font.color.rgb
                current_bold = run.font.bold
                base.set_run_font(
                    run,
                    size=None,
                    bold=current_bold,
                    color=str(current_color) if current_color is not None else None,
                )
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        current_color = run.font.color.rgb
                        current_bold = run.font.bold
                        base.set_run_font(
                            run,
                            size=TABLE_SIZE,
                            bold=current_bold,
                            color=str(current_color) if current_color is not None else None,
                        )

    document.save(OUTPUT)
    restore_preserved_package_parts(SOURCE, OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    try:
        build()
    except Exception as exc:
        print(f"build failed: {exc}", file=sys.stderr)
        raise
