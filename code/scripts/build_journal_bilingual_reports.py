#!/usr/bin/env python3
"""Build journal-styled Chinese and English report counterparts.

The Chinese edition is derived from the reviewed Morandi report.  The English
edition is generated from the translated DOCX supplied through ``--english-source``
and then receives the identical table/caption/layout treatment.
"""
from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[2]
SOURCE_ZH = PROJECT / "output/docx/DASE7506_Project_1_TECH_Report_莫兰迪图表版.docx"
OUTPUT_ZH = PROJECT / "output/docx/DASE7506_Project_1_TECH_Report_中文期刊格式版.docx"
OUTPUT_EN = PROJECT / "output/docx/DASE7506_Project_1_TECH_Report_English_Journal_Style.docx"

MORANDI_SLATE = "7E919D"
MORANDI_HEADER = "E8ECE7"
MORANDI_ALT = "F6F7F4"
TEXT = "252B2F"
LIGHT = "B8C0C3"


def _border(parent, edge: str, *, val: str, size: int = 0, color: str = "auto") -> None:
    tag = parent.find(qn(f"w:{edge}"))
    if tag is None:
        tag = OxmlElement(f"w:{edge}")
        parent.append(tag)
    tag.set(qn("w:val"), val)
    tag.set(qn("w:sz"), str(size))
    tag.set(qn("w:space"), "0")
    tag.set(qn("w:color"), color)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def _set_cell_margins(cell, top=80, start=85, bottom=80, end=85) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        marker = OxmlElement("w:tblHeader")
        marker.set(qn("w:val"), "true")
        tr_pr.append(marker)


def _format_three_line_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    _border(borders, "top", val="single", size=12, color=MORANDI_SLATE)
    _border(borders, "bottom", val="single", size=12, color=MORANDI_SLATE)
    for edge in ("left", "right", "insideH", "insideV"):
        _border(borders, edge, val="nil")

    # The middle rule belongs to the header row; no other body rules are used.
    header = table.rows[0]
    _set_repeat_header(header)
    for cell in header.cells:
        _shade(cell, MORANDI_HEADER)
        tc_pr = cell._tc.get_or_add_tcPr()
        cell_borders = tc_pr.find(qn("w:tcBorders"))
        if cell_borders is None:
            cell_borders = OxmlElement("w:tcBorders")
            tc_pr.append(cell_borders)
        _border(cell_borders, "bottom", val="single", size=7, color=MORANDI_SLATE)

    for r_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell)
            if r_index > 0:
                _shade(cell, "FFFFFF" if r_index % 2 else MORANDI_ALT)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor.from_string(TEXT)
                    run.font.bold = r_index == 0


FIGURE_RE = re.compile(r"^(?:扩展图\s*E\d+|图\s*B\d+|图\s*\d+)\s*[.:：]?\s*(.*)$", re.I)
TABLE_RE = re.compile(r"^(?:扩展表\s*E\d+|表\s*\d+)\s*[.:：]?\s*(.*)$", re.I)
FIGURE_EN_RE = re.compile(r"^(?:Extended\s+Figure\s+E?\d+|Figure\s+B?\d+|Fig\.\s*\d+)\s*[.:]?\s*(.*)$", re.I)
TABLE_EN_RE = re.compile(r"^(?:Extended\s+Table\s+E?\d+|Table\s+\d+)\s*[.:]?\s*(.*)$", re.I)


def _replace_caption(paragraph, label: str, rest: str) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = False
    lead = paragraph.add_run(label)
    lead.bold = True
    lead.font.size = Pt(9)
    lead.font.color.rgb = RGBColor.from_string(TEXT)
    body = paragraph.add_run(rest)
    body.font.size = Pt(9)
    body.font.color.rgb = RGBColor.from_string(TEXT)


def _format_captions(doc: Document, *, english: bool) -> tuple[int, int]:
    fig_no = 0
    table_no = 0
    figure_pattern = FIGURE_EN_RE if english else FIGURE_RE
    table_pattern = TABLE_EN_RE if english else TABLE_RE
    paragraphs = doc.paragraphs
    figure_caption_nodes = {
        paragraphs[i + 1]._p
        for i, p in enumerate(paragraphs[:-1])
        if p._p.xpath('.//w:drawing')
    }
    for paragraph in paragraphs:
        text = " ".join(paragraph.text.split())
        m = figure_pattern.match(text)
        if m and paragraph._p in figure_caption_nodes:
            fig_no += 1
            rest = m.group(1).strip()
            _replace_caption(paragraph, f"Fig. {fig_no}. ", rest)
            continue
        if english and m and paragraph._p not in figure_caption_nodes:
            # A prior formatting pass may have labeled an ordinary paragraph
            # beginning with "Figure ...".  It is explanatory prose, not a caption.
            paragraph.text = m.group(1).strip()
            continue
        m = table_pattern.match(text)
        if m:
            table_no += 1
            rest = m.group(1).strip()
            _replace_caption(paragraph, f"Table {table_no}. ", rest)
            paragraph.paragraph_format.keep_with_next = True
    return fig_no, table_no


ENGLISH_FIXES = {
    "Table 19. RoPE tied / untied 配对": "Table 19. Paired comparison of tied and untied RoPE models",
    "8 RoPE 三 seed 与最佳 checkpoint": "8 RoPE Across Three Seeds and Best-Checkpoint Selection",
    "8.2 training loss 下降为何不能选末尾权重": "8.2 Why a Lower Training Loss Does Not Justify Selecting the Final Checkpoint",
    "B.1 v1 与初次 RoPE 长训练": "B.1 v1 and the Initial Long RoPE Run",
    "BPE-2048；context 256；独立 causal 窗口；CPU FP32": "BPE-2048; context 256; independent causal windows; CPU FP32",
    "固定 tokenizer 与 evaluator": "Fixed tokenizer and evaluator",
    "总 NLL ÷ ln(2) ÷ UTF-8 bytes": "Total NLL / ln(2) / UTF-8 bytes",
    "exp（总 NLL ÷ target 数）": "exp(total NLL / number of targets)",
    "小型 RMS + SwiGLU": "Small RMSNorm + SwiGLU",
    "v1 不使用 dropout": "Use no dropout in v1",
    "小型 RMS+SwiGLU": "Small RMSNorm + SwiGLU",
    "总 NLL（nats）": "Total NLL (nats)",
    "seed 17：learned tied 1.820432，untied 1.786678，RoPE tied 1.754858；RoPE 双 seed 均值 1.748021": "Seed 17: learned tied 1.820432, untied 1.786678, and RoPE tied 1.754858; the two-seed mean for RoPE was 1.748021",
    "相对 learned tied 降低": "Reduction relative to learned tied",
    "EMA 0.99，同 run live": "EMA 0.99 versus the live weights from the same run",
    "MTP k=1，匹配 untied 配方": "MTP k=1 with the matched untied recipe",
    "固定 BPE-2048、context 256 与 evaluator": "Fixed BPE-2048, context 256, and evaluator",
    "相同 processed targets 的比较": "Comparison at the same number of processed targets",
    "第 2、6 节与 freeze manifest": "Sections 2 and 6 and the freeze manifest",
    "训练成本、seed 与 checkpoint ancestry": "Training cost, seeds, and checkpoint ancestry",
    "5× CPU、4 GiB、64 MiB 限制": "5x CPU, 4 GiB, and 64 MiB limits",
    "03 小型 RMS+SwiGLU / LR .001": "03 Small RMSNorm+SwiGLU / LR .001",
    "04 小型 RMS+SwiGLU / LR .002": "04 Small RMSNorm+SwiGLU / LR .002",
    "13 原始 baseline 设置": "13 Original baseline setup",
    "RoPE 三 seed validation": "RoPE three-seed validation",
    "Reasoning about assets": "Inference assets",
    "Waiting time for the same flight ÷ baseline time": "Candidate runtime / baseline runtime on the same machine",
    "Courses are baseline": "Course baseline",
    "ICE": "GELU",
    "1,873 GiB": "1.873 GiB",
    "1,948 GiB": "1.948 GiB",
    "2,875 times": "2.875x",
}


def _cleanup_english(doc: Document) -> None:
    """Repair untranslated fragments and a few known machine-translation errors."""
    containers = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                containers.extend(cell.paragraphs)
    for paragraph in containers:
        old = paragraph.text
        new = ENGLISH_FIXES.get(old, old)
        for source, replacement in ENGLISH_FIXES.items():
            if source in new:
                new = new.replace(source, replacement)
        if new != old:
            paragraph.text = new


def _format_document(source: Path, target: Path, *, english: bool) -> tuple[int, int]:
    shutil.copy2(source, target)
    doc = Document(target)
    if english:
        _cleanup_english(doc)
    for table in doc.tables:
        _format_three_line_table(table)
    figures, tables = _format_captions(doc, english=english)
    # Journal-like restraint: black headings, centered title, consistent body color.
    for style_name in ("Title", "Heading 1", "Heading 2"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.color.rgb = RGBColor(0, 0, 0)
    for p in doc.paragraphs:
        if p.style and p.style.name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title = (
        "DASE7506 Project 1 Technical Report English Journal Style"
        if english else "DASE7506 Project 1 技术报告 中文期刊格式版"
    )
    doc.save(target)
    return figures, tables


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--english-source", type=Path)
    args = parser.parse_args()
    OUTPUT_ZH.parent.mkdir(parents=True, exist_ok=True)
    zf, zt = _format_document(SOURCE_ZH, OUTPUT_ZH, english=False)
    print(f"Chinese: {OUTPUT_ZH} ({zf} figures, {zt} table captions, {len(Document(OUTPUT_ZH).tables)} tables)")
    if args.english_source:
        ef, et = _format_document(args.english_source, OUTPUT_EN, english=True)
        print(f"English: {OUTPUT_EN} ({ef} figures, {et} table captions, {len(Document(OUTPUT_EN).tables)} tables)")


if __name__ == "__main__":
    main()
